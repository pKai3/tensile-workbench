"""Standalone numerical library for Tensile Workbench v16.

No batch entry point, graph definitions, file discovery on import, or
dependency on plot_tensile scripts. Calculations preserved from v15.
All output writes require explicit calls with destination paths.
"""

import sys

import csv

import re

import itertools

import shutil

import os

import tempfile

from datetime import datetime

from pathlib import Path

import numpy as np

import pandas as pd

import matplotlib.pyplot as plt

from tensile_properties_v18 import prepared_curve, specimen_properties

NAME_LOOKUP = {}  # Display labels come from the saved project.

EXCEL_DECIMAL_PLACES = 3

LANDMARK_ERROR_BARS = True

PLOT_FIGSIZE = (9, 6)

PLOT_DPI = 600

LANDMARK_POINTS_PER_STAGE = 400

LANDMARK_PREBREAK_MARGIN_PERCENT = 0.05

LANDMARK_YIELD_FIT_FRACTIONS = (0.20, 0.50)

LANDMARK_YIELD_R2_WARNING = 0.98

POINTS = 600

TENSILE_MEAN_TAIL = {
    "enabled": True,
    "setback_percent": 0.20,       # retreat before earliest fracture-drop onset
    "fit_window_percent": 3.0,    # fit only this final measured strain window
    "polynomial_order": 2,        # 1 = linear, 2 = quadratic
    "max_extension_ratio": 2.0,   # maximum extrapolation / fitted strain span
    "adaptive_fit": False,       # widen fit / permit a peak in a rising tail
    "fracture_slope_fraction": 0.30,  # lower = detect slower fracture drops
}

YOUNGS_MODULUS_MPA = 115000.0

WH_MIN_PLASTIC_STRAIN_PERCENT = 0.2

WH_POINTS = 800

WH_DESPIKE_ENABLED = True

WH_DESPIKE_WINDOW_PERCENT = 0.30

WH_DESPIKE_SIGMA = 4.0

WH_MEDIAN_FILTER_ENABLED = True

WH_MEDIAN_FILTER_WINDOW_PERCENT = 0.30

WH_SMOOTH_ENABLED = True

WH_SMOOTH_WINDOW_PERCENT = 0.75

WH_POLY_ORDER = 2

WH_XLIM = (1.0, 10.0)

WH_YLIM = (0, 4500)

FORCE_COLS = [
    "Load", "Force", "Load (N)", "Force (N)", "Force(N)", "Load(N)",
    "Axial Force", "Axial Force (N)"
]

STRESS_COLS = [
    "Stress", "Engineering Stress", "True Stress", "Stress (MPa)",
    "Engineering Stress (MPa)", "True Stress (MPa)",
    "Stress (Pa)", "Engineering Stress (Pa)", "True Stress (Pa)",
    "Tensile stress"
]

STRAIN_COLS = [
    "Strain", "Engineering Strain", "True Strain", "Strain (%)",
    "Engineering Strain (%)", "True Strain (%)",
    "Axial Strain", "Axial Strain (%)",
    "Tensile strain (Strain 1)", "Strain 1"
]

UNIT_ROW_PATTERN = re.compile(r"^\s*\([^)]+\)\s*$")

def get_display_name(folder_name: str, overrides=None) -> str:
    if overrides and folder_name in overrides:
        return overrides[folder_name]
    return NAME_LOOKUP.get(folder_name, folder_name)

def find_first_col(df, candidates):
    for c in candidates:
        if c in df.columns:
            return c
    lower_map = {col.lower(): col for col in df.columns}
    for c in candidates:
        if c.lower() in lower_map:
            return lower_map[c.lower()]
    return None

def detect_units_row(df):
    if df.shape[0] == 0:
        return False, {}

    row0 = df.iloc[0]
    matches, units = 0, {}
    for col, val in row0.items():
        if isinstance(val, str) and UNIT_ROW_PATTERN.match(val):
            matches += 1
            units[col] = val.strip()[1:-1]

    has_units = matches >= max(2, int(0.4 * len(df.columns)))
    return has_units, (units if has_units else {})

def coerce_numeric(df, skip_first_row=False):
    if skip_first_row and df.shape[0] > 0:
        df = df.iloc[1:].reset_index(drop=True)
    for c in df.columns:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    return df

def as_float_series(s):
    return pd.to_numeric(s, errors="coerce")

def compute_stress_mpa(df, units_map):
    """
    Use stress column if present.
    If only force exists, skip the file because geometry-based stress calculation
    has been removed.
    """
    stress_col = find_first_col(df, STRESS_COLS)
    if stress_col:
        stress = as_float_series(df[stress_col])
        unit = (units_map.get(stress_col) or "").lower()

        if unit == "mpa":
            return stress.rename("Stress (MPa)")
        if unit == "pa":
            return (stress / 1e6).rename("Stress (MPa)")

        med = np.nanmedian(stress.values) if np.isfinite(stress.values).any() else np.nan
        if np.isfinite(med) and med > 1e4:
            return (stress / 1e6).rename("Stress (MPa)")

        return stress.rename("Stress (MPa)")

    force_col = find_first_col(df, FORCE_COLS)
    if force_col:
        print(f"[WARN] Found force column '{force_col}' but no stress column; skipping.")
    return None

def compute_strain_percent(df, units_map):
    """
    Return strain as percent (%).

    Handles:
      - Explicit percent columns (kept as-is)
      - Fraction-style strain columns (converted ×100 using heuristic)

    Does NOT compute strain from extension/displacement.
    """
    strain_col = find_first_col(df, STRAIN_COLS)
    if strain_col is not None:
        s = as_float_series(df[strain_col])
        unit = (units_map.get(strain_col) or "").strip().lower()

        if unit in ["%", "percent", "percent (%)"]:
            return s.rename("Strain (%)")

        try:
            p50 = float(np.nanpercentile(s, 50))
            p90 = float(np.nanpercentile(s, 90))
        except Exception:
            p50, p90 = np.nan, np.nan

        if (np.isfinite(p50) and p50 <= 2.0) and (np.isfinite(p90) and p90 <= 5.0):
            return (s * 100.0).rename("Strain (%)")

        return s.rename("Strain (%)")

    return None

def compute_toughness_safe(strain_percent, stress_mpa):
    """
    Toughness for one specimen in MJ/m^3.
    """
    s_pct = np.asarray(strain_percent, dtype=float)
    st = np.asarray(stress_mpa, dtype=float)

    mask = np.isfinite(s_pct) & np.isfinite(st)
    if mask.sum() < 2:
        return np.nan

    return np.trapezoid(st[mask], s_pct[mask] / 100.0)

def summarize_toughness_stats(models, plot_spec):
    """Compare synthetic curve integrals with the mean specimen integral.

    Stress (MPa) times dimensionless engineering strain gives MJ/m^3. Integrate
    the exact sampled curves used for display, regardless of plot axis limits.
    A close area match is a diagnostic, not proof of pointwise curve accuracy.
    Do not report a full-curve percentage difference for an incomplete mean.
    """
    rows = []
    for g, model in sorted(models.items()):
        curves = model["curves"]
        T = [compute_toughness_safe(s_pct, st) for (s_pct, st) in curves]
        T = [t for t in T if np.isfinite(t)]

        n_valid = len(T)
        mean_T = float(np.mean(T)) if n_valid > 0 else np.nan
        sd_T = float(np.std(T, ddof=1)) if n_valid > 1 else np.nan

        pointwise, landmark = model["pointwise"], model["landmark"]
        pointwise_area = compute_toughness_safe(pointwise["x"], pointwise["y"]) if pointwise else np.nan
        landmark_area = compute_toughness_safe(landmark["x"], landmark["y"]) if landmark else np.nan
        complete = bool(pointwise and pointwise["complete"])
        def difference(area, valid=True):
            return 100*(area-mean_T)/mean_T if valid and np.isfinite(area) and np.isfinite(mean_T) and mean_T != 0 else np.nan
        rows.append({
            "Graph": plot_spec["name"],
            "Group": g,
            "Display Name": get_display_name(g, plot_spec.get("name_overrides")),
            "n_curves": len(curves),
            "n_valid": n_valid,
            "Toughness Mean (MJ/m^3)": mean_T,
            "Toughness SD (MJ/m^3)": sd_T,
            "Toughness Min (MJ/m^3)": float(np.min(T)) if n_valid > 0 else np.nan,
            "Toughness Max (MJ/m^3)": float(np.max(T)) if n_valid > 0 else np.nan,
            "Pointwise Curve Area (MJ/m^3)": pointwise_area,
            "Pointwise Difference (%)": difference(pointwise_area, complete),
            "Pointwise Measured Area (MJ/m^3)": compute_toughness_safe(pointwise["measured_x"], pointwise["measured_y"]) if pointwise else np.nan,
            "Pointwise Fitted Tail Area (MJ/m^3)": (compute_toughness_safe(pointwise["predicted_x"], pointwise["predicted_y"]) if len(pointwise["predicted_x"]) else 0.0) if pointwise else np.nan,
            "Pointwise Coverage": "complete to mean terminal strain" if complete else ("partial" if pointwise else "unavailable"),
            "Pointwise Start Strain (%)": float(pointwise["x"][0]) if pointwise else np.nan,
            "Pointwise End Strain (%)": float(pointwise["x"][-1]) if pointwise else np.nan,
            "Pointwise Diagnostic": model["pointwise_diagnostic"],
            "Landmark Curve Area (MJ/m^3)": landmark_area,
            "Landmark Difference (%)": difference(landmark_area),
            "Landmark Start Strain (%)": float(landmark["x"][0]) if landmark else np.nan,
            "Landmark End Strain (%)": float(landmark["x"][-1]) if landmark else np.nan,
            "Landmark Diagnostic": model["landmark_diagnostic"] or "stage-aligned; pre-break endpoint mapped to mean terminal strain",
            "Mean Failure Elongation (%)": model["mean_failure"],
        })
    return rows

def compute_uniform_properties(strain_percent, engineering_stress_mpa):
    """
    Return engineering UTS and uniform elongation for one specimen.

    Uniform elongation is the engineering strain at maximum engineering stress.
    The first maximum is used if the measured curve contains an exact plateau.
    """
    s_pct = np.asarray(strain_percent, dtype=float)
    stress = np.asarray(engineering_stress_mpa, dtype=float)
    mask = np.isfinite(s_pct) & np.isfinite(stress)
    s_pct = s_pct[mask]
    stress = stress[mask]

    if len(s_pct) < 3:
        return np.nan, np.nan

    i_uts = int(np.nanargmax(stress))
    return float(stress[i_uts]), float(s_pct[i_uts])

def local_polynomial_derivative(x, y, window=31, order=2):
    """
    Smooth and differentiate y(x) with a centred local polynomial fit.

    This is a small NumPy-only equivalent of a Savitzky-Golay derivative and
    avoids amplifying point-to-point noise as strongly as a raw gradient.
    """
    x = np.asarray(x, dtype=float)
    y = np.asarray(y, dtype=float)
    n = len(x)
    if n < max(order + 2, 5):
        return np.full(n, np.nan)

    w = max(int(window), order + 2)
    if w % 2 == 0:
        w += 1
    w = min(w, n if n % 2 == 1 else n - 1)
    if w < order + 2:
        return np.full(n, np.nan)

    half = w // 2
    derivative = np.full(n, np.nan)

    for i in range(n):
        lo = max(0, i - half)
        hi = min(n, i + half + 1)
        if hi - lo < w:
            if lo == 0:
                hi = min(n, w)
            elif hi == n:
                lo = max(0, n - w)

        dx = x[lo:hi] - x[i]
        yy = y[lo:hi]
        finite = np.isfinite(dx) & np.isfinite(yy)
        if finite.sum() < order + 2:
            continue

        coeff = np.polyfit(dx[finite], yy[finite], order)
        derivative[i] = coeff[-2]

    return derivative

def strain_window_to_points(strain_percent, window_percent, minimum_points=5):
    """Convert a window width in strain percentage points to an odd point count."""
    x = np.asarray(strain_percent, dtype=float)
    steps = np.diff(x)
    steps = steps[np.isfinite(steps) & (steps > 0)]
    if len(steps) == 0 or not np.isfinite(window_percent) or window_percent <= 0:
        return int(minimum_points) | 1

    points = max(int(round(window_percent / float(np.median(steps)))), minimum_points)
    if points % 2 == 0:
        points += 1
    return points

def hampel_filter(values, window, sigma=4.0):
    """
    Replace isolated outliers using a rolling-median/MAD Hampel filter.

    End regions that cannot support a centred window are left unchanged. The
    input array is not modified. Returns the filtered values and replacement
    count.
    """
    y = np.asarray(values, dtype=float)
    filtered = y.copy()
    n = len(y)
    w = max(int(window), 5)
    if w % 2 == 0:
        w += 1
    if n < w:
        return filtered, 0

    half = w // 2
    replacements = 0
    for i in range(half, n - half):
        local = y[i - half:i + half + 1]
        local = local[np.isfinite(local)]
        if len(local) < 5 or not np.isfinite(y[i]):
            continue

        median = float(np.median(local))
        mad = float(np.median(np.abs(local - median)))
        robust_sigma = 1.4826 * mad
        if robust_sigma <= 0:
            continue

        if abs(y[i] - median) > float(sigma) * robust_sigma:
            filtered[i] = median
            replacements += 1

    return filtered, replacements

def rolling_median_filter(values, window):
    """Apply a centred rolling median while leaving unsupported ends unchanged."""
    y = np.asarray(values, dtype=float)
    filtered = y.copy()
    n = len(y)
    w = max(int(window), 5)
    if w % 2 == 0:
        w += 1
    if n < w:
        return filtered

    half = w // 2
    for i in range(half, n - half):
        local = y[i - half:i + half + 1]
        local = local[np.isfinite(local)]
        if len(local) >= 5:
            filtered[i] = float(np.median(local))
    return filtered

def compute_work_hardening_curve(
    strain_percent,
    engineering_stress_mpa,
    youngs_modulus_mpa=YOUNGS_MODULUS_MPA,
    min_plastic_strain_percent=WH_MIN_PLASTIC_STRAIN_PERCENT,
    num_points=WH_POINTS,
    despike_enabled=WH_DESPIKE_ENABLED,
    despike_window_percent=WH_DESPIKE_WINDOW_PERCENT,
    despike_sigma=WH_DESPIKE_SIGMA,
    median_filter_enabled=WH_MEDIAN_FILTER_ENABLED,
    median_filter_window_percent=WH_MEDIAN_FILTER_WINDOW_PERCENT,
    smooth_enabled=WH_SMOOTH_ENABLED,
    smooth_window_percent=WH_SMOOTH_WINDOW_PERCENT,
    poly_order=WH_POLY_ORDER,
):
    """
    Calculate work-hardening rate up to the onset of necking.

    Returns true plastic strain (%) and theta (MPa), where
    theta = d(true stress) / d(true plastic strain).
    """
    s_pct = np.asarray(strain_percent, dtype=float)
    stress_eng = np.asarray(engineering_stress_mpa, dtype=float)
    mask = np.isfinite(s_pct) & np.isfinite(stress_eng)
    s_pct = s_pct[mask]
    stress_eng = stress_eng[mask]

    if len(s_pct) < 10 or not np.isfinite(youngs_modulus_mpa) or youngs_modulus_mpa <= 0:
        return None

    order_idx = np.argsort(s_pct)
    s_pct = s_pct[order_idx]
    stress_eng = stress_eng[order_idx]
    s_pct, unique_idx = np.unique(s_pct, return_index=True)
    stress_eng = stress_eng[unique_idx]

    if len(s_pct) < 10:
        return None

    # Maximum engineering stress marks the end of uniform deformation.
    i_uts = int(np.nanargmax(stress_eng))
    if i_uts < 8:
        return None
    s_pct = s_pct[:i_uts + 1]
    stress_eng = stress_eng[:i_uts + 1]

    strain_eng = s_pct / 100.0
    stress_true = stress_eng * (1.0 + strain_eng)
    strain_true = np.log1p(strain_eng)
    strain_plastic = strain_true - stress_true / youngs_modulus_mpa

    min_eps_p = min_plastic_strain_percent / 100.0
    valid = (
        np.isfinite(strain_plastic)
        & np.isfinite(stress_true)
        & (strain_plastic >= min_eps_p)
    )
    strain_plastic = strain_plastic[valid]
    stress_true = stress_true[valid]

    if len(strain_plastic) < 10:
        return None

    order_idx = np.argsort(strain_plastic)
    strain_plastic = strain_plastic[order_idx]
    stress_true = stress_true[order_idx]
    strain_plastic, unique_idx = np.unique(strain_plastic, return_index=True)
    stress_true = stress_true[unique_idx]

    if len(strain_plastic) < 10 or strain_plastic[-1] <= strain_plastic[0]:
        return None

    grid = np.linspace(strain_plastic[0], strain_plastic[-1], int(num_points))
    stress_grid = np.interp(grid, strain_plastic, stress_true)
    grid_percent = grid * 100.0

    if despike_enabled:
        despike_window = strain_window_to_points(
            grid_percent,
            despike_window_percent,
            minimum_points=5,
        )
        stress_grid, _ = hampel_filter(
            stress_grid,
            window=despike_window,
            sigma=despike_sigma,
        )

    if median_filter_enabled:
        median_window = strain_window_to_points(
            grid_percent,
            median_filter_window_percent,
            minimum_points=5,
        )
        stress_grid = rolling_median_filter(stress_grid, median_window)

    if smooth_enabled:
        smooth_window = strain_window_to_points(
            grid_percent,
            smooth_window_percent,
            minimum_points=poly_order + 3,
        )
        theta = local_polynomial_derivative(
            grid,
            stress_grid,
            window=smooth_window,
            order=poly_order,
        )
    else:
        theta = np.gradient(stress_grid, grid)

    return grid_percent, theta, stress_grid

def require_excel_export():
    """Fail before a run writes any outputs if its Excel dependency is absent."""
    try:
        import openpyxl
    except ImportError as error:
        raise RuntimeError(
            'Excel output requires openpyxl. Install it in this Python environment:\n'
            f'  "{sys.executable}" -m pip install openpyxl'
        ) from error
    if not isinstance(EXCEL_DECIMAL_PLACES, int) or not 0 <= EXCEL_DECIMAL_PLACES <= 15:
        raise ValueError("EXCEL_DECIMAL_PLACES must be an integer from 0 to 15.")
    return openpyxl

def write_excel_table(frame, output_path, sheet_name="Results"):
    """Export a plain numeric table, formatting display without rounding data.

    One former CSV becomes one workbook. Missing numeric results stay blank,
    strings remain literal text, and existing files are replaced only after a
    complete temporary workbook has been saved successfully.
    """
    openpyxl = require_excel_export()
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    output_path = Path(output_path)
    if output_path.suffix.lower() != ".xlsx":
        raise ValueError(f"Excel output requires an .xlsx filename: {output_path}")
    if len(frame) + 1 > 1048576 or len(frame.columns) > 16384:
        raise ValueError(f"Table exceeds Excel worksheet limits: {output_path}")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    number_format = "0" + ("." + "0" * EXCEL_DECIMAL_PLACES if EXCEL_DECIMAL_PLACES else "")
    count_columns = {"n", "n_curves", "n_valid", "count"}
    body_font = Font(name="Arial", size=10)
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="24445C")
    numeric_alignment = Alignment(horizontal="right", vertical="center")
    text_alignment = Alignment(horizontal="left", vertical="center")
    headers = [str(column) for column in frame.columns]
    widths = [min(32, max(14, len(header) + 2)) for header in headers]

    for col, header in enumerate(headers, 1):
        cell = sheet.cell(1, col, header)
        cell.data_type = "s"
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 45

    for row_index, row in enumerate(frame.itertuples(index=False, name=None), 2):
        for col_index, value in enumerate(row, 1):
            if isinstance(value, np.generic):
                value = value.item()
            if value is None or pd.isna(value):
                value = None
            elif isinstance(value, float) and not np.isfinite(value):
                value = str(value)  # preserve +/-inf as an explicit nonnumeric result
            cell = sheet.cell(row_index, col_index, value)
            cell.font = body_font
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.number_format = "0" if headers[col_index-1].lower() in count_columns else number_format
                cell.alignment = numeric_alignment
            else:
                cell.alignment = text_alignment
                if isinstance(value, str):
                    # A sample/path/label beginning '=' must never become a formula.
                    cell.data_type = "s"
                    widths[col_index-1] = min(60, max(widths[col_index-1], len(value) + 2))

    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    if headers:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(frame)+1}"
    else:
        sheet["A1"] = "No data available"
        sheet["A1"].font = body_font
        sheet.column_dimensions["A"].width = 24

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(prefix=f".{output_path.stem}_", suffix=".xlsx",
                                         dir=output_path.parent, delete=False) as temporary:
            temporary_path = Path(temporary.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
    except PermissionError as error:
        raise PermissionError(f"Cannot write {output_path}. Close it in Excel and check folder permissions.") from error
    finally:
        workbook.close()
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()
    print(f"[OK] Wrote {output_path}")

def trim_nonnegative_percent(strain_percent, stress):
    mask = (~strain_percent.isna()) & (~stress.isna())
    strain_percent, stress = strain_percent[mask], stress[mask]

    if len(strain_percent) == 0:
        return strain_percent.iloc[0:0], stress.iloc[0:0]

    cond = (strain_percent.values >= 0) & (stress.values >= 0)
    if not np.any(cond):
        return strain_percent.iloc[0:0], stress.iloc[0:0]

    first = np.argmax(cond)
    strain_percent, stress = strain_percent.iloc[first:], stress.iloc[first:]

    s, st = strain_percent.values, stress.values
    keep = np.concatenate(([True], np.diff(s) >= 0))
    return pd.Series(s[keep]), pd.Series(st[keep])

def read_tensile_csv_table(csv_path):
    """Find the measurement header, ignoring any preceding results tables.

    Match complete supported column names, not fragments such as "Tensile
    stress at Maximum Force" in an Instron summary. Keep the units row for the
    existing conversion logic. The source file is never modified.
    """
    strain_names = {name.lower() for name in STRAIN_COLS}
    load_names = {name.lower() for name in STRESS_COLS + FORCE_COLS}

    for encoding in ("utf-8-sig", "cp1252"):
        for sep in (",", ";", "\t"):
            try:
                with open(csv_path, encoding=encoding, newline="") as source:
                    reader = csv.reader(source, delimiter=sep)
                    header_line = 0
                    for row in reader:
                        names = {value.strip().lower() for value in row}
                        if names & strain_names and names & load_names:
                            break
                        # Count physical lines, including embedded newlines in
                        # quoted summary fields.
                        header_line = reader.line_num
                    else:
                        continue

                    # Rewind to the actual header. Passing the remaining stream
                    # avoids pandas skiprows counting quoted multiline records
                    # differently from physical lines in the summary section.
                    source.seek(0)
                    for _ in range(header_line):
                        source.readline()
                    frame = pd.read_csv(source, sep=sep)
                frame.columns = [str(column).strip() for column in frame.columns]
                frame.attrs["measurement_header_line"] = header_line + 1
                return frame
            except (OSError, UnicodeError, csv.Error, pd.errors.ParserError,
                    pd.errors.EmptyDataError):
                continue
    return None


def load_and_prepare_curve(csv_path):
    """Load the complete recorded curve; display-only trimming happens later."""
    df_raw = read_tensile_csv_table(csv_path)

    if df_raw is None:
        return None

    has_units, units_map = detect_units_row(df_raw)
    df = coerce_numeric(df_raw, skip_first_row=has_units)

    stress = compute_stress_mpa(df, units_map if has_units else {})
    strain_pct = compute_strain_percent(df, units_map if has_units else {})

    if stress is None or strain_pct is None:
        return None

    strain_pct, stress = trim_nonnegative_percent(strain_pct, stress)
    if len(strain_pct) < 5:
        return None

    order = np.argsort(strain_pct.values)
    s_sorted = strain_pct.values[order]
    st_sorted = stress.values[order]

    return s_sorted, st_sorted

def find_sample_groups(root: Path):
    groups = {}
    for p in sorted([
        d for d in root.iterdir()
        if d.is_dir() and not d.name.startswith("!")
    ]):
        csvs = [f for f in sorted(p.rglob("*.csv")) if not f.name.startswith("!")]
        if csvs:
            groups[p.name] = csvs
    return groups

def average_group_curves(curves, num_points=600):
    """Fixed-population mean over the interval shared by every valid curve."""
    if not curves:
        return None, None, None, None

    valid_curves = [c for c in curves if c is not None and len(c[0]) > 1 and len(c[1]) > 1]
    if not valid_curves:
        return None, None, None, None

    starts = [np.nanmin(c[0]) for c in valid_curves]
    ends = [np.nanmax(c[0]) for c in valid_curves]

    grid_min = max(starts)
    grid_max = min(ends)

    if not np.isfinite(grid_min) or not np.isfinite(grid_max) or grid_max <= grid_min:
        return None, None, None, None

    grid = np.linspace(grid_min, grid_max, num_points)

    mats = []
    for s_pct, st in valid_curves:
        uniq_s, uniq_idx = np.unique(s_pct, return_index=True)
        uniq_st = st[uniq_idx]
        if len(uniq_s) < 2:
            continue
        mats.append(np.interp(grid, uniq_s, uniq_st, left=np.nan, right=np.nan))

    if not mats:
        return None, None, None, None

    M = np.vstack(mats)
    counts = np.sum(np.isfinite(M), axis=0)

    valid_cols = counts == len(valid_curves)
    if not np.any(valid_cols):
        return None, None, None, None

    grid = grid[valid_cols]
    M = M[:, valid_cols]
    counts = counts[valid_cols]

    mean = np.nanmean(M, axis=0)
    std = np.nanstd(M, axis=0)

    if np.all(~np.isfinite(mean)):
        return None, None, None, None

    return grid, mean, std, counts

def tensile_fracture_onset(strain, stress, slope_fraction=0.30):
    """Locate an abrupt terminal drop, not the ordinary post-UTS decline.

    A uniform strain grid makes detection independent of acquisition frequency.
    If no abrupt drop is found, use the recorded terminal strain. This is a
    plotting heuristic, not a replacement for the instrument's break metric.
    """
    strain, stress = np.asarray(strain), np.asarray(stress)
    valid = np.isfinite(strain) & np.isfinite(stress)
    x, idx = np.unique(strain[valid], return_index=True)
    y = stress[valid][idx]
    if len(x) < 20 or x[-1] <= x[0]:
        return float(x[-1]) if len(x) else np.nan
    grid = np.linspace(x[0], x[-1], min(5000, max(100, int((x[-1]-x[0])/0.01)+1)))
    values = np.interp(grid, x, y)
    slope = np.diff(values) / np.diff(grid)
    peak = int(np.argmax(values))
    end_region = max(peak + 1, int(0.65 * len(slope)))
    baseline = slope[peak:end_region]
    if len(baseline) < 10:
        if slope_fraction >= 0.30:
            return float(x[-1])
        # Abrupt failures close to UTS have too little post-peak baseline.
        median, mad = 0.0, 0.0
    else:
        median = float(np.median(baseline))
        mad = float(1.4826*np.median(np.abs(baseline-median)))
    limit = max(8*abs(median), 8*mad, slope_fraction*float(np.max(values)))
    drops = np.flatnonzero(slope[end_region:] < median-limit)
    # Ignore isolated signal pings that recover shortly after the sharp drop.
    lookahead = max(2, int(np.ceil(0.15/(grid[1]-grid[0]))))
    for drop in drops:
        i = end_region+drop
        later = values[i+1:min(len(values), i+1+lookahead)]
        if len(later) and later[-1] < values[i]-0.01*np.max(values):
            return float(grid[i])
    return float(x[-1])

def build_tensile_mean_tail(curves, points, settings):
    """Return measured x/y, predicted x/y, and diagnostic for plots/area checks."""
    order = settings["polynomial_order"]
    window = float(settings["fit_window_percent"])
    setback = float(settings["setback_percent"])
    ratio = float(settings["max_extension_ratio"])
    adaptive = settings.get("adaptive_fit", False)
    slope_fraction = float(settings.get("fracture_slope_fraction", 0.30))
    if not np.isfinite(slope_fraction) or slope_fraction <= 0:
        raise ValueError("fracture_slope_fraction must be positive.")
    if order not in (1, 2) or not np.isfinite([window,setback,ratio]).all() or window <= 0 or setback < 0 or ratio <= 0:
        raise ValueError("Tensile mean-tail settings require order 1/2, positive window/ratio and nonnegative setback.")
    clean = []
    for x, y in curves:
        x, y = np.asarray(x), np.asarray(y)
        valid = np.isfinite(x) & np.isfinite(y)
        x, idx = np.unique(x[valid], return_index=True)
        y = y[valid][idx]
        if len(x) >= 2:
            clean.append((x, y))
    empty = np.array([], dtype=float)
    if not clean:
        return empty, empty, empty, empty, "no valid curves"
    if len(clean) < 2:
        x, y = clean[0]
        return x, y, empty, empty, "one specimen; no mean-tail estimate"
    start = max(x[0] for x, _ in clean)
    end = min(tensile_fracture_onset(x, y, slope_fraction) for x, y in clean)-setback
    target = float(np.mean([x[-1] for x, _ in clean]))
    if end <= start:
        return empty, empty, empty, empty, "no common pre-fracture interval"
    x = np.linspace(start, end, max(30, points))
    y = np.mean([np.interp(x, sx, sy) for sx, sy in clean], axis=0)
    if adaptive:
        return fit_adaptive_tensile_tail(x, y, target, points, settings)
    # Fit only post-peak data, preventing a local fit across the tensile maximum.
    fit_start = max(end-window, x[int(np.argmax(y))])
    fit = x >= fit_start
    span = end-x[fit][0]
    if fit.sum() < 8 or span <= 0:
        return x, y, empty, empty, "insufficient post-peak data for tail fit"
    if target-end > ratio*span:
        return x, y, empty, empty, "extension exceeds max_extension_ratio; estimate omitted"
    # Anchor the polynomial exactly to the last measured average value.
    u = (x[fit]-end)/span
    px = np.linspace(end, target, max(30, int(points*(target-end)/(end-start))))
    pu = (px-end)/span
    for degree in range(order, 0, -1):
        matrix = np.column_stack([u**i for i in range(1, degree+1)])
        coeff = np.linalg.lstsq(matrix, y[fit]-y[-1], rcond=None)[0]
        py = y[-1]+sum(c*pu**i for i,c in enumerate(coeff, 1))
        gradient = sum(i*c*pu**(i-1) for i,c in enumerate(coeff, 1))
        # Reject upward-turning or non-positive post-necking extrapolations.
        if np.isfinite(py).all() and np.all(py > 0) and np.all(gradient <= 1e-8):
            return x, y, px, py, f"order {degree}; join {end:.2f}%, mean end {target:.2f}%"
    return x, y, empty, empty, "polynomial failed tail-shape checks; estimate omitted"

def fit_adaptive_tensile_tail(x, y, target, points, settings):
    """Anchored local polynomial, with wider data support for variable ductility.

    Prefer the requested window; widen only to satisfy the extension limit or
    shape checks. A rising join may turn through one maximum, but cannot grow
    indefinitely. Failed fits remain omitted, not clipped into plausible shapes.
    """
    empty = np.array([], dtype=float)
    end = x[-1]
    extension = target-end
    window = max(float(settings["fit_window_percent"]), 1.05*extension/float(settings["max_extension_ratio"]))
    # Stay clear of elastic loading and the initial yielding transition.
    plastic_start = max(x[0], 1.0)
    px = np.linspace(end, target, max(30, int(points*extension/(end-x[0]))))
    for width in (window, 1.5*window, 2*window):
        fit = x >= max(end-width, plastic_start)
        span = end-x[fit][0] if fit.any() else 0
        if fit.sum()<8 or span<=0 or extension>float(settings["max_extension_ratio"])*span:
            continue
        u=(x[fit]-end)/span
        pu=(px-end)/span
        # Robust near-join trend determines whether an initial rise is allowed.
        recent = x >= max(end-min(1.0, span/3), x[fit][0])
        trend = np.polyfit(x[recent]-end, y[recent], 1)[0]
        for degree in range(settings["polynomial_order"], 0, -1):
            matrix=np.column_stack([u**i for i in range(1, degree+1)])
            coeff=np.linalg.lstsq(matrix,y[fit]-y[-1],rcond=None)[0]
            predicted=y[-1]+sum(c*pu**i for i,c in enumerate(coeff,1))
            gradient=sum(i*c*pu**(i-1) for i,c in enumerate(coeff,1))
            residual=(y[-1]+matrix@coeff)-y[fit]
            # A poor description of the measured tail is not safe to project.
            if np.sqrt(np.mean(residual**2))>0.015*np.max(y):
                continue
            valid=np.isfinite(predicted).all() and np.all(predicted>0)
            if trend<=0:
                valid = valid and np.all(gradient<=1e-8)
            else:
                valid = valid and degree==2 and coeff[1]<0 and gradient[0]>0 and gradient[-1]<=0
                # Limit the projected rise to 10% above the measured maximum.
                valid = valid and np.max(predicted)<=1.10*np.max(y)
            if valid:
                return x,y,px,predicted,f"adaptive order {degree}, fit {span:.2f} pp; join {end:.2f}%, mean end {target:.2f}%"
    return x,y,empty,empty,"adaptive fit lacks support or fails shape checks; estimate omitted"

def draw_landmark_points(ax, landmark, color, error_bars, comparison=False):
    """Draw specimen scatter at fixed mean landmarks, not curve-fit uncertainty."""
    for point in landmark["point_statistics"]:
        # With bars off, retain the v11 marker behavior in comparison plots.
        if comparison and not error_bars and point["name"] != "failure":
            continue
        if error_bars and landmark["n"] > 1:
            ax.errorbar(point["x"], point["y"], xerr=point["x_sd"],
                        yerr=point["y_sd"], fmt=point["marker"],
                        color=color, ecolor=color, ms=point["marker_size"],
                        elinewidth=.9, capsize=3, capthick=.9, linestyle="none",
                        label="_nolegend_", zorder=4)
        else:
            ax.plot(point["x"], point["y"], point["marker"], color=color,
                    ms=point["marker_size"], label="_nolegend_", zorder=4)

def finish_plot(figure, out_path, preview=False):
    """Return a live figure for previews; save/close only for batch calls."""
    if not preview:
        figure.savefig(out_path, dpi=PLOT_DPI)
        plt.close(figure)
        print(f"[OK] Saved {out_path}")
    return figure

def render_average_plot(models, records, out_path, family, show_individual,
                        color_map, xlim, ylim, title, name_overrides=None,
                        landmark_error_bars=None, preview=False):
    """Render cached means so paired/comparison plots and area stats agree."""
    figure, ax = plt.subplots(figsize=PLOT_FIGSIZE)
    comparison = family == "comparison"
    if landmark_error_bars is None:
        landmark_error_bars = LANDMARK_ERROR_BARS
    plotted = False
    projected = False
    for group, model in sorted(models.items()):
        pointwise, landmark = model["pointwise"], model["landmark"]
        if comparison and (pointwise is None or landmark is None):
            print(f"[WARN] Comparison omits {group}: both mean forms are required.")
            continue
        if not comparison and model[family] is None:
            continue
        color = color_map[group]
        label = f"{get_display_name(group, name_overrides)} (n={len(records[group])})"
        if show_individual:
            for record in records[group]:
                ax.plot(record["strain_pct"], record["stress_mpa"], color=color, alpha=.18, lw=.7)
        if family in ("pointwise", "comparison"):
            ax.plot(pointwise["measured_x"], pointwise["measured_y"], color=color, label=label, lw=1.8)
            if len(pointwise["predicted_x"]):
                ax.plot(pointwise["predicted_x"], pointwise["predicted_y"], color=color, ls=":", lw=1.8)
                ax.plot(pointwise["x"][-1], pointwise["y"][-1], "x", color=color, ms=6)
                projected = True
        if family in ("landmark", "comparison"):
            ax.plot(landmark["x"], landmark["y"], color=color,
                    ls="--" if comparison else "-", lw=1.8,
                    label=None if comparison else label)
            draw_landmark_points(ax, landmark, color, landmark_error_bars, comparison)
        plotted = True
    if not plotted:
        plt.close(figure)
        print(f"[WARN] No valid {family} curves for {out_path}")
        return
    ax.set(title=title, xlabel="Engineering strain (%)", ylabel="Engineering stress (MPa)", xlim=xlim, ylim=ylim)
    ax.legend()
    ax.grid(True, ls="--", alpha=.6)
    if comparison:
        note = "Solid: pointwise; dotted: fitted tail; dashed: landmark-aligned."
    elif family == "landmark":
        note = "○ Mean YS   △ Mean UTS   × Mean failure elongation"
    else:
        note = "Solid: common-population pointwise mean.\nDotted: fitted tail; × at mean failure elongation." if projected else "Common-population pointwise mean; no fitted tail."
    show_bar_note = family in ("landmark", "comparison") and landmark_error_bars
    if show_bar_note:
        note += "\nLandmark bars: ±1 SD (n ≥ 2); endpoint stress is pre-break."
    figure.text(.5, .01, note, ha="center")
    figure.tight_layout(rect=(0, .08, 1, 1))
    return finish_plot(figure, out_path, preview)

def failure_elongation_percent(strain_percent):
    """Return the terminal engineering strain of a prepared tensile curve."""
    strain = np.asarray(strain_percent, dtype=float)
    finite = strain[np.isfinite(strain)]
    return float(np.nanmax(finite)) if len(finite) else np.nan

def choose_representative_curve(group_name, records, override=None):
    """
    Select one specimen from a group for the representative-curves plot.

    The automatic choice minimizes absolute distance from the group's mean
    terminal strain. An override may match a CSV stem, filename, or full path.
    """
    valid_records = [
        record for record in records
        if np.isfinite(record.get("failure_elongation_pct", np.nan))
    ]
    if not valid_records:
        return None

    if override is not None:
        requested = str(override).strip().casefold()
        matches = []
        for record in valid_records:
            source_path = Path(record["source_file"])
            candidates = {
                str(record["sample"]).casefold(),
                source_path.name.casefold(),
                source_path.stem.casefold(),
                str(source_path).casefold(),
            }
            if requested in candidates:
                matches.append(record)

        if len(matches) == 1:
            return matches[0]
        if len(matches) > 1:
            print(
                f"[WARN] Representative override '{override}' for '{group_name}' "
                "matched more than one specimen; using the first match."
            )
            return matches[0]
        print(
            f"[WARN] Representative override '{override}' for '{group_name}' "
            "was not found; using automatic selection."
        )

    mean_failure = float(np.mean([
        record["failure_elongation_pct"] for record in valid_records
    ]))
    return min(
        valid_records,
        key=lambda record: (
            abs(record["failure_elongation_pct"] - mean_failure),
            str(record["sample"]).casefold(),
        ),
    )

def render_representative_tensile_plot(
    groups_records,
    out_path,
    color_map,
    xlim,
    ylim,
    name_overrides=None,
    representative_overrides=None,
    title="Representative stress–strain curves",
    show_individual=False,
    preview=False,
):
    """Plot one representative engineering stress-strain curve per group."""
    plt.figure(figsize=PLOT_FIGSIZE)
    plotted = False
    representative_overrides = representative_overrides or {}

    for group_name in sorted(groups_records.keys()):
        representative = choose_representative_curve(
            group_name,
            groups_records[group_name],
            override=representative_overrides.get(group_name),
        )
        if representative is None:
            continue

        if show_individual:
            for record in groups_records[group_name]:
                plt.plot(record["strain_pct"], record["stress_mpa"],
                         lw=.7, alpha=.18, color=color_map[group_name])
        plt.plot(
            representative["strain_pct"],
            representative["stress_mpa"],
            lw=1.6,
            color=color_map[group_name],
            label=get_display_name(group_name, name_overrides),
        )
        print(
            f"[REP] {out_path.stem}: {group_name} -> "
            f"{representative['sample']} "
            f"({representative['failure_elongation_pct']:.2f}% failure elongation)"
        )
        plotted = True

    if not plotted:
        print(
            f"[WARN] Nothing plotted for {out_path} (no valid representative curves).",
            file=sys.stderr,
        )
        plt.close()
        return

    plt.xlabel("Strain (%)")
    plt.ylabel("Stress (MPa)")
    plt.title(title)
    plt.legend()
    plt.xlim(xlim)
    plt.ylim(ylim)
    plt.grid(True, linestyle="--", alpha=0.6)
    plt.tight_layout()
    return finish_plot(plt.gcf(), out_path, preview)

def render_work_hardening_plot(
    groups_curves,
    out_path,
    color_map,
    show_individual=False,
    xlim=WH_XLIM,
    ylim=WH_YLIM,
    name_overrides=None,
    title="Work-hardening rate before necking",
    youngs_modulus_mpa=YOUNGS_MODULUS_MPA,
    youngs_modulus_overrides=None,
    min_plastic_strain_percent=WH_MIN_PLASTIC_STRAIN_PERCENT,
    num_points=WH_POINTS,
    wh_filter_settings=None,
    preview=False,
):
    """
    Plot theta = d(true stress)/d(true plastic strain) up to uniform elongation.
    """
    if not np.isfinite(youngs_modulus_mpa) or youngs_modulus_mpa <= 0:
        raise ValueError("youngs_modulus_mpa must be finite and positive (MPa).")
    if not np.isfinite(min_plastic_strain_percent) or min_plastic_strain_percent < 0:
        raise ValueError("wh_min_plastic_strain_percent must be finite and non-negative.")
    if isinstance(num_points, bool) or not isinstance(num_points, (int, np.integer)) or num_points < 10:
        raise ValueError("wh_points must be an integer of at least 10.")
    modulus_overrides = youngs_modulus_overrides or {}
    for group_name, modulus in modulus_overrides.items():
        if not np.isfinite(modulus) or modulus <= 0:
            raise ValueError(f"Invalid modulus for {group_name}: use a positive value in MPa.")
        if group_name not in groups_curves:
            print(f"[WARN] Modulus override '{group_name}' matches no selected data folder.",
                  file=sys.stderr)
    plt.figure(figsize=PLOT_FIGSIZE)
    plotted = False
    all_rates = []
    x_max = 0.0

    for group_name in sorted(groups_curves.keys()):
        wh_curves = []
        for s_pct, stress_eng in groups_curves[group_name]:
            result = compute_work_hardening_curve(
                s_pct, stress_eng,
                youngs_modulus_mpa=modulus_overrides.get(group_name, youngs_modulus_mpa),
                min_plastic_strain_percent=min_plastic_strain_percent,
                num_points=num_points,
                **(wh_filter_settings or {}),
            )
            if result is None:
                continue
            eps_p_pct, theta, _ = result
            valid = np.isfinite(eps_p_pct) & np.isfinite(theta)
            if valid.sum() < 5:
                continue
            eps_p_pct = eps_p_pct[valid]
            theta = theta[valid]
            wh_curves.append((eps_p_pct, theta))
            all_rates.extend(theta[theta >= 0].tolist())
            x_max = max(x_max, float(np.nanmax(eps_p_pct)))

        if not wh_curves:
            continue

        averaged = average_group_curves(
            wh_curves,
            num_points=num_points,
        )
        if averaged[0] is None:
            continue

        grid_pct, mean_theta, _, _ = averaged
        color = color_map[group_name]

        if show_individual:
            for eps_p_pct, theta in wh_curves:
                plt.plot(eps_p_pct, theta, alpha=0.18, lw=0.7, color=color)

        plt.plot(
            grid_pct,
            mean_theta,
            lw=1.8,
            color=color,
            label=f"{get_display_name(group_name, name_overrides)} (n={len(wh_curves)})",
        )
        plotted = True

    if not plotted:
        print(f"[WARN] Nothing plotted for {out_path} (no valid work-hardening curves).",
              file=sys.stderr)
        plt.close()
        return

    plt.xlabel("True plastic strain (%)")
    plt.ylabel(r"Work-hardening rate, $\Theta$ (MPa)")
    plt.title(title)
    if xlim is not None:
        plt.xlim(xlim)
    else:
        plt.xlim(0, x_max * 1.03 if x_max > 0 else 1)

    if ylim is not None:
        plt.ylim(ylim)
    elif all_rates:
        robust_max = float(np.nanpercentile(np.asarray(all_rates), 99.0))
        if np.isfinite(robust_max) and robust_max > 0:
            plt.ylim(0, robust_max * 1.10)

    plt.legend()
    plt.grid(True, linestyle="--", alpha=0.6)
    plt.tight_layout()
    return finish_plot(plt.gcf(), out_path, preview)

def compute_landmark_work_hardening_curve(landmark, **settings):
    """Differentiate the aligned mean only through its UTS landmark.

    This is a synthetic, stage-aligned diagnostic, not a population-average
    hardening rate. Do not include the post-UTS section or a fitted tensile tail.
    The original WH conversion and filters are reused without modification.
    """
    if landmark is None:
        return None
    end = float(landmark["knots"][2])
    x, y = landmark["x"], landmark["y"]
    before_peak = x < end
    loading_x = np.append(x[before_peak], end)
    loading_y = np.append(y[before_peak], landmark["mean_uts"])
    return compute_work_hardening_curve(loading_x, loading_y, **settings)

def render_landmark_work_hardening_plot(
    models, groups_curves, out_path, color_map, show_individual=False,
    xlim=WH_XLIM, ylim=WH_YLIM, name_overrides=None,
    title="Landmark-derived work-hardening response",
    youngs_modulus_mpa=YOUNGS_MODULUS_MPA, youngs_modulus_overrides=None,
    min_plastic_strain_percent=WH_MIN_PLASTIC_STRAIN_PERCENT, num_points=WH_POINTS,
    wh_filter_settings=None, preview=False,
):
    """Plot the derivative of each landmark mean; optional faint specimen rates."""
    if not np.isfinite(youngs_modulus_mpa) or youngs_modulus_mpa <= 0:
        raise ValueError("youngs_modulus_mpa must be finite and positive (MPa).")
    if not np.isfinite(min_plastic_strain_percent) or min_plastic_strain_percent < 0:
        raise ValueError("wh_min_plastic_strain_percent must be finite and non-negative.")
    if isinstance(num_points, bool) or not isinstance(num_points, (int, np.integer)) or num_points < 10:
        raise ValueError("wh_points must be an integer of at least 10.")
    overrides = youngs_modulus_overrides or {}
    for group, modulus in overrides.items():
        if not np.isfinite(modulus) or modulus <= 0:
            raise ValueError(f"Invalid modulus for {group}: use a positive value in MPa.")
        if group not in groups_curves:
            print(f"[WARN] Modulus override '{group}' matches no selected data folder.", file=sys.stderr)
    figure, ax = plt.subplots(figsize=PLOT_FIGSIZE)
    plotted = False
    all_rates = []
    x_max = 0.0
    for group, model in sorted(models.items()):
        settings = {
            **(wh_filter_settings or {}),
            "youngs_modulus_mpa": overrides.get(group, youngs_modulus_mpa),
            "min_plastic_strain_percent": min_plastic_strain_percent,
            "num_points": num_points,
        }
        landmark = model["landmark"]
        result = compute_landmark_work_hardening_curve(landmark, **settings)
        if result is None:
            print(f"[WH LANDMARK WARNING] {Path(out_path).stem} / {group}: unavailable; plot omitted.")
            continue
        x, theta, _ = result
        valid = np.isfinite(x) & np.isfinite(theta)
        if valid.sum() < 5:
            continue
        x, theta = x[valid], theta[valid]
        color = color_map[group]
        if show_individual:
            for strain, stress in groups_curves[group]:
                individual = compute_work_hardening_curve(strain, stress, **settings)
                if individual is not None:
                    ix, iy, _ = individual
                    iv = np.isfinite(ix) & np.isfinite(iy)
                    ax.plot(ix[iv], iy[iv], color=color, alpha=.18, lw=.7)
        ax.plot(x, theta, color=color, lw=1.8,
                label=f"{get_display_name(group, name_overrides)} (n={landmark['n']})")
        all_rates.extend(theta[theta >= 0].tolist())
        x_max = max(x_max, float(x[-1]))
        plotted = True
        print(f"[WH LANDMARK] {Path(out_path).stem} / {group}: derivative of aligned mean; "
              f"mean uniform elongation={landmark['knots'][2]:.3f}%; "
              f"ends at true plastic strain={x[-1]:.3f}%; "
              f"E={settings['youngs_modulus_mpa']:.1f} MPa.")
    if not plotted:
        plt.close(figure)
        print(f"[WARN] No valid landmark-derived work-hardening curves for {out_path}")
        return
    ax.set(title=title, xlabel="True plastic strain (%)",
           ylabel=r"Work-hardening rate, $\Theta$ (MPa)")
    ax.set_xlim(xlim if xlim is not None else (0, x_max * 1.03 if x_max > 0 else 1))
    if ylim is not None:
        ax.set_ylim(ylim)
    elif all_rates:
        robust_max = float(np.nanpercentile(all_rates, 99.0))
        if np.isfinite(robust_max) and robust_max > 0:
            ax.set_ylim(0, robust_max * 1.10)
    ax.legend()
    ax.grid(True, linestyle="--", alpha=.6)
    note = "Derivative of landmark-aligned mean, not mean specimen hardening rate."
    note += "\nEnds at mean uniform elongation."
    if show_individual:
        note += " Faint lines: individual specimens."
    figure.text(.5, .01, note, ha="center")
    figure.tight_layout(rect=(0, .08, 1, 1))
    return finish_plot(figure, out_path, preview)

def group_number(group_name):
    """Extract the leading B-number, e.g. B005-2 -> 5."""
    match = re.match(r"^[Bb]0*(\d+)", str(group_name))
    return int(match.group(1)) if match else None

def matches_group_selector(group_name, selector):
    """Match an exact group or a hyphenated family such as B11 -> B11-A1."""
    name = str(group_name).casefold()
    requested = str(selector).strip().casefold()
    return name == requested or name.startswith(requested + "-")

def select_plot_groups(groups_curves, ranges=None, group_selectors=None):
    """Select the union of inclusive numeric ranges and explicit group selectors."""
    ranges = ranges or []
    group_selectors = group_selectors or []
    selected = {}
    for group_name, curves in groups_curves.items():
        number = group_number(group_name)
        range_match = (
            number is not None
            and any(int(start) <= number <= int(end) for start, end in ranges)
        )
        explicit_match = any(
            matches_group_selector(group_name, selector)
            for selector in group_selectors
        )
        if range_match or explicit_match:
            selected[group_name] = curves
    return selected

def safe_filename(text):
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", str(text)).strip("_.")
    return cleaned or "plot_set"

def landmark_specimen(record, fit_fractions=LANDMARK_YIELD_FIT_FRACTIONS,
                      points_per_stage=None):
    """Three measured sections, each mapped to progress 0..1 (no extrapolation)."""
    x, y = prepared_curve(record)
    if len(x) < 20:
        raise ValueError("fewer than 20 valid strain points")
    peak = int(np.argmax(y))
    raw_uts, peak_x = float(y[peak]), float(x[peak])
    if raw_uts <= 0 or peak < 3 or peak >= len(x)-2:
        raise ValueError("missing resolved loading or post-UTS segment")
    uts,end_x=raw_uts,float(x[-1])
    details = specimen_properties(record, fit_fractions, LANDMARK_YIELD_R2_WARNING)
    if details['Yield status'] != 'resolved':
        raise ValueError(details['Notes'])
    yield_x, ys = details['Yield strain (%)'], details['Yield (MPa)']
    # Remove terminal fracture unloading from shape, but retain the raw terminal
    # failure strain as the destination landmark. The end stress is pre-break.
    onset=tensile_fracture_onset(x,y,slope_fraction=0.10)
    clean_end=min(end_x,onset-LANDMARK_PREBREAK_MARGIN_PERCENT)
    if clean_end <= peak_x:
        # A very short post-peak segment still needs two measured positions.
        clean_end=min(end_x,float(x[peak+1]))
    if not x[0]<yield_x<peak_x<clean_end:
        raise ValueError("landmarks are not ordered start < yield < UTS < pre-break")
    source_knots=np.array([x[0],yield_x,peak_x,clean_end])
    target_knots=np.array([x[0],yield_x,peak_x,end_x])
    count = LANDMARK_POINTS_PER_STAGE if points_per_stage is None else points_per_stage
    if isinstance(count, bool) or not isinstance(count, (int, np.integer)) or count < 3:
        raise ValueError("Landmark points per stage must be an integer of at least 3")
    progress=np.linspace(0,1,count)
    stages=[np.interp(a+progress*(b-a),x,y) for a,b in zip(source_knots[:-1],source_knots[1:])]
    stages[0][-1]=stages[1][0]=ys
    stages[1][-1]=stages[2][0]=uts
    details.update({"Yield (MPa)":ys,"Yield strain (%)":yield_x,"UTS (MPa)":uts,
                    "Uniform elongation (%)":peak_x,"Failure elongation (%)":end_x,
                    "Pre-break source strain (%)":clean_end,"Pre-break stress (MPa)":float(stages[2][-1])})
    return target_knots,np.array(stages),details

def average_landmark_shapes(items):
    """Equal specimen weights at equal progress within each deformation stage."""
    knots=np.mean([item[0] for item in items],axis=0)
    stages=np.mean([item[1] for item in items],axis=0)
    progress=np.linspace(0,1,stages.shape[1])
    xs=[];ys=[]
    for j in range(3):
        cut=slice(None) if j==0 else slice(1,None)
        xs.extend((knots[j]+progress*(knots[j+1]-knots[j]))[cut])
        ys.extend(stages[j][cut])
    return np.array(xs),np.array(ys),knots

def prepare_average_curves(plot_spec, selected_records, *,
                           landmark_points_per_stage=None, pointwise_points=None):
    """Compute each mean once for plots, comparisons, exports, and integration."""
    landmark_points = (LANDMARK_POINTS_PER_STAGE if landmark_points_per_stage is None
                       else landmark_points_per_stage)
    points = POINTS if pointwise_points is None else pointwise_points
    if isinstance(landmark_points, bool) or not isinstance(landmark_points, (int, np.integer)) or landmark_points<3:
        raise ValueError("LANDMARK_POINTS_PER_STAGE must be at least 3")
    if isinstance(points, bool) or not isinstance(points, (int, np.integer)) or points<10:
        raise ValueError("Pointwise points must be an integer of at least 10")
    audit, summaries, models = [], [], {}
    tail_settings = dict(TENSILE_MEAN_TAIL)
    tail_settings.update(plot_spec.get("tensile_mean_tail", {}))
    for group,records in sorted(selected_records.items()):
        if not records:
            continue
        curves = [(r["strain_pct"], r["stress_mpa"]) for r in records]
        target = float(np.mean([failure_elongation_percent(x) for x, _ in curves]))
        empty = np.array([], dtype=float)
        if tail_settings["enabled"]:
            mx, my, px, py, diagnostic = build_tensile_mean_tail(curves, points, tail_settings)
        else:
            result = average_group_curves(curves, points)
            mx, my = (result[0], result[1]) if result[0] is not None else (empty, empty)
            px, py, diagnostic = empty, empty, "tail disabled; common interval only"
        pointwise = None
        if len(mx) >= 2:
            # The first predicted point duplicates the measured join. Include it
            # only once so integration and the displayed joined curve agree.
            full_x = np.concatenate((mx, px[1:]))
            full_y = np.concatenate((my, py[1:]))
            pointwise = {"x": full_x, "y": full_y, "measured_x": mx,
                         "measured_y": my, "predicted_x": px, "predicted_y": py,
                         "complete": bool(np.isclose(full_x[-1], target, rtol=0, atol=1e-6))}
        print(f"[TAIL] {group}: {diagnostic}")
        models[group] = {"pointwise": pointwise, "landmark": None,
                         "mean_failure": target, "pointwise_diagnostic": diagnostic,
                         "landmark_diagnostic": "", "curves": curves}
        fit_fractions=plot_spec.get("landmark_yield_fit_fractions",LANDMARK_YIELD_FIT_FRACTIONS)
        items=[];rejected=[]
        for record in records:
            try:
                item=landmark_specimen(record,fit_fractions,points_per_stage=landmark_points)
                if item[2]["Notes"]:
                    print(f"[LANDMARK WARNING] {group}/{record['sample']}: {item[2]['Notes']}")
                items.append(item); audit.append({"Group":group,"Status":"included",**item[2]})
            except ValueError as error:
                rejected.append(record["sample"])
                audit.append({"Group":group,"Sample":record["sample"],"Raw source":record["source_file"],
                              "Status":"unresolved", "Notes":str(error)})
        # Do not silently change the population to make a synthetic mean work.
        if rejected or not items:
            models[group]["landmark_diagnostic"] = f"unresolved specimens: {', '.join(rejected)}"
            print(f"[LANDMARK WARNING] {group}: plot omitted; unresolved specimens {rejected}")
            continue
        x,y,knots=average_landmark_shapes(items)
        models[group]["landmark"] = {"x": x, "y": y, "knots": knots,
                                      "mean_yield": float(np.mean([i[2]["Yield (MPa)"] for i in items])),
                                      "mean_uts": float(np.mean([i[2]["UTS (MPa)"] for i in items]))}
        summary={"Group":group,"n":len(items)}
        for key in ("Yield (MPa)","Yield strain (%)","UTS (MPa)","Uniform elongation (%)","Failure elongation (%)","Pre-break stress (MPa)"):
            summary["Mean "+key]=float(np.mean([i[2][key] for i in items]))
            summary["SD "+key]=float(np.std([i[2][key] for i in items], ddof=1)) if len(items)>1 else np.nan
        # Derive error bars from original per-specimen landmark properties,
        # never from the aligned/resampled mean curve.
        models[group]["landmark"]["n"] = len(items)
        models[group]["landmark"]["point_statistics"] = [
            {"name": name, "x": summary["Mean "+x_key], "y": summary["Mean "+y_key],
             "x_sd": summary["SD "+x_key], "y_sd": summary["SD "+y_key],
             "marker": marker, "marker_size": size}
            for name,x_key,y_key,marker,size in (
                ("yield", "Yield strain (%)", "Yield (MPa)", "o", 4),
                ("UTS", "Uniform elongation (%)", "UTS (MPa)", "^", 5),
                ("failure", "Failure elongation (%)", "Pre-break stress (MPa)", "x", 6))]
        summaries.append(summary)
        print(f"[LANDMARK] {group}: n={len(items)}, YS={summary['Mean Yield (MPa)']:.2f}, UTS={summary['Mean UTS (MPa)']:.2f}, EL={summary['Mean Failure elongation (%)']:.2f}%")
    return models, audit, summaries

def compute_axes(groups_curves):
    x_max, y_max = 0.0, 0.0
    for curves in groups_curves.values():
        for s_pct, st in curves:
            if s_pct.size:
                x_max = max(x_max, float(np.nanmax(s_pct)))
            if st.size:
                y_max = max(y_max, float(np.nanmax(st)))
    x_max = x_max if x_max > 0 else 1.0
    y_max = y_max if y_max > 0 else 1.0
    return (0, x_max * 1.05), (0, y_max * 1.05)
