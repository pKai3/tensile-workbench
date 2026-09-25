"""Two-workbook exports; numeric, wide curve columns and no hidden plot methods."""
from datetime import datetime
from pathlib import Path
import json
import os
import tempfile
import numpy as np
import pandas as pd
from tensile_gauge import group_basis_label, group_policy
from tensile_fit import fit_display_places
from tensile_fracture import FRACTURE_METHOD


def write_workbook(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    path = Path(path)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, frame in sheets.items():
        sheet = workbook.create_sheet(name[:31])
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = 'B2' if name in ('Summary', 'Specimens', 'Audit') else 'A2'
        if frame.empty and not len(frame.columns):
            sheet.append(['No data available'])
            continue
        if len(frame) >= 1048576 or len(frame.columns) > 16384:
            raise ValueError(f'{name} exceeds Excel worksheet limits')
        headers = [str(c) for c in frame.columns]
        sheet.append(headers)
        for row in frame.itertuples(index=False, name=None):
            values = []
            for v in row:
                if isinstance(v, np.generic):
                    v = v.item()
                if v is None or (not isinstance(v, (list, dict)) and pd.isna(v)):
                    v = None
                elif isinstance(v, float) and not np.isfinite(v):
                    v = None
                elif isinstance(v, (dict, list, tuple)):
                    v = json.dumps(v)
                values.append(v)
            sheet.append(values)
        for cells in sheet.iter_rows():
            for cell in cells:
                cell.font = Font(name='Arial', size=10)
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # literal names/paths, never formulas
                    if headers[cell.column - 1] == 'Checks':
                        cell.alignment = Alignment(wrap_text=True, vertical='top')
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    header = headers[cell.column - 1].lower()
                    integer = (header.endswith(' n') or header in ('n', 'included n', 'available n', 'pairs with checks')
                               or 'row (1-based)' in header)
                    cell.number_format = '0' if integer else '0.' + '0' * fit_display_places(header)
        for cell in sheet[1]:
            cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='24445C')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[1].height = 45
        for i, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(i)].width = min(38, max(18, len(header) * .65))
            if header == 'Checks':
                sheet.column_dimensions[get_column_letter(i)].width = 65
        sheet.auto_filter.ref = sheet.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, prefix='.' + path.stem, suffix='.xlsx', delete=False) as f:
            temporary = Path(f.name)
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        workbook.close()
        if temporary is not None and temporary.exists():
            temporary.unlink()


def flatten_info(metadata):
    rows = [{'Setting': 'Exported at', 'Value': datetime.now().astimezone().isoformat()}]
    def visit(key, value):
        if isinstance(value, dict):
            for k, v in value.items():
                visit(f'{key}.{k}' if key else k, v)
        elif isinstance(value, (list, tuple)):
            for i, v in enumerate(value):
                visit(f'{key}[{i}]', v)
        else:
            rows.append({'Setting': key, 'Value': value})
    visit('', metadata)
    return pd.DataFrame(rows)


def export_results(frames, path, metadata):
    from tensile_tables import specimen_export_frame, summary_wide, SPECIMEN_DETAIL_FIELDS
    from tensile_properties import EL_SOURCE_FIELDS
    samples = frames['tensile_samples']
    checks = frames['specimen_diagnostics'].copy()
    comparison = frames['instron_comparison']
    # Use exactly the UI's numeric schema, not its formatted strings or the
    # internal acquisition dictionaries. Search filters never change exports.
    specimens = specimen_export_frame(samples, comparison)
    summary = summary_wide(frames['tensile_summary_details'])
    # Provenance and acquisition bookkeeping belong only in the audit sheet.
    # Keep stable specimen IDs for joins; never depend on display row order.
    visible_fields = {'Included', 'Exclusion Reason', 'Inclusion Scope', 'Checks',
                      '0.2% Offset Yield Strength (MPa)', 'UTS (MPa)', 'Uniform Elongation (%)',
                      'Failure Elongation (%)', 'Measured failure elongation (%)', 'Reconstructed EL (%)',
                      'Toughness (MJ/m^3)', 'Fitted E (GPa)',
                      *(field for _, field in EL_SOURCE_FIELDS),
                      *(field for _, field in SPECIMEN_DETAIL_FIELDS)}
    if not samples.empty:
        source = samples.set_index('Specimen ID')
        for field in samples:
            if field not in visible_fields and field not in checks:
                checks[field] = checks['Specimen ID'].map(source[field])
        # These lengths are useful for auditing the gauge calculation even
        # though the corresponding display columns remain in Specimens.
        for field in ('AVE dot spacing (mm)', 'Target gauge length (mm)'):
            checks[field] = checks['Specimen ID'].map(source[field]) if field in source else np.nan
    if 'Estimated toughness (MJ/m^3)' in checks:
        applied = checks['Gauge correction status'].eq('Applied') & checks['Elongation basis'].eq('Estimated standard gauge')
        checks['Estimated toughness (MJ/m^3)'] = checks['Estimated toughness (MJ/m^3)'].where(applied)
        checks = checks.rename(columns={'Estimated toughness (MJ/m^3)': 'Reconstructed toughness (MJ/m^3)'})
    if 'Elongation basis' in checks:
        checks['Elongation basis'] = checks['Elongation basis'].replace({'Estimated standard gauge': 'Reconstructed gauge'})
    lookup = {'0.2% yield strength': '0.2% Offset Yield Strength (MPa)',
              'UTS': 'UTS (MPa)', 'Uniform elongation': 'Uniform Elongation (%)',
              'Fitted elastic modulus': 'Fitted E (GPa)'}
    for label, field in lookup.items():
        if comparison.empty or field not in samples:
            continue
        match = comparison[comparison['Property'].str.casefold() == label.casefold()].drop_duplicates('Specimen ID').set_index('Specimen ID')
        for key in ('Difference (calc - Instron)', 'Difference (%)'):
            checks[label + ' · ' + key] = checks['Specimen ID'].map(match[key])
    if 'Override Definition' in checks:
        def description(value):
            try:
                obj = json.loads(value) if value else {}
                return '; '.join(f'{key}: {obj[key]}' for key in ('mode', 'strain_bounds', 'endpoints') if key in obj)
            except (ValueError, TypeError):
                return str(value)
        checks['Override Definition'] = checks['Override Definition'].map(description)
    # Shared IDs are join keys; avoid repeating the main property values here.
    checks = checks.drop(columns=[c for c in (visible_fields | {'Elastic fit R2'}) if c in checks
                                 and c not in ('AVE dot spacing (mm)', 'Target gauge length (mm)')])
    metadata = {**metadata, 'definitions': {
        'calculated_basis': 'Tensile plots/statistics use the selected CSV endpoint, or its gauge reconstruction. A valid saved manual row overrides automatic selection. Toughness integrates through the chosen measurement. Automatic selection uses a supported sudden drop, otherwise the first below-10%-of-peak crossing. Truncated exports may supply review-flagged sudden-drop or progressive terminal-unloading estimates; absence of 2%/10% confirmation does not automatically suppress EL. No unconditional final-reading fallback. Pre-peak properties and WH use measured data. No Instron substitution.',
        'elongation_columns': 'Instron is imported Strain 1 at break. Calc is the selected CSV endpoint, without a strain setback. Reconstruct applies the gauge model to that same endpoint and is blank unless enabled and available. Audit preserves the selected criterion, review reason, original rows/time, load-loss evidence and settings, plus last valid/max retained strain.',
        'instron_comparison': 'Differences compare measured calculations with original Instron values, never reconstructed values.',
        'checks': 'Failure/review messages appear in the Specimens Checks column. Identity is verified by name/export row and original CSV UTS agreement within printed-digit rounding tolerance. EL is not an identity check and differences from Instron break EL are not matching failures. No automatic exclusion or numerical rematching.',
        'table_layout': 'Summary and Specimens share the viewer\'s property names, order, analysis basis and blank-value rules. Uniform elongation and tensile toughness retain Calc and Instron columns; unsupported or missing imports remain blank. Elongation has Instron, Calc and Reconstruct columns. Summary mean, SD and n are separate numeric columns. Specimens adds a stable Specimen ID to join Audit. Audit contains source, fit, gauge and comparison details, not an additional user table.',
        'selection': 'Global specimen data-use modes apply unless this graph explicitly overrides them. After-UTS AVE failure retains YS/UTS/E/uniform EL and pre-peak WH; after-yield failure retains YS/UTS/E; unreliable strain throughout retains UTS only. All partial modes exclude full tensile shape, failure EL and full-test toughness, including corresponding Instron properties. Specimens retains raw values and records exclusions; summary and comparison statistics use eligible values only. Hard-ignored ! files never load.',
        'landmark_populations': 'Trusted eligible curves supply stage shapes, remapped between each property\'s eligible group mean coordinates. Strength and elongation may have different n. Error bars are coordinate-specific specimen SD, not fit uncertainty. No shape or invalid ordered mean landmarks means no curve. Pointwise curves use eligible shape specimens only; strength-EL scatter uses paired eligible specimens.',
        'gauge': 'Strain 1 gauge length is the AVE-measured initial dot spacing. Target and enabled state are per sample group within this graph.',
        'estimate': 'Post-peak reconstruction assumes both gauge intervals capture the localisation. Longer targets are allowed with warnings. Not a standards-compliant measurement.',
        'gauge_formula': 'r = L_dots / L_target; pre-peak strain unchanged; post-peak strain_est = strain_u + r * (strain_measured - strain_u). Failure EL_est = EL_u + r * (EL_measured - EL_u). All strains are total engineering strain in percent; stress unchanged throughout.',
        'failure_endpoint': FRACTURE_METHOD + '. Original force sequence (engineering stress proxy if force absent), no smoothing or strain sorting. Adjacent-reading sudden-drop rule: loss >5 times the preceding change magnitude, confirmed below 2% of peak (ISO 6892-1:2019 informative A.3.6). Substantial unconfirmed single-step drops and supported multi-reading collapses can be algorithmically detected without implying ISO confirmation. A separate heuristic detects sustained multi-reading collapse: at least 5% event-level loss, rapid unloading relative to local background, and no sustained recovery. It selects the rapid-change leading edge, or a sharper transition within an accelerating lead-in, without requiring a fraction of the later maximum unloading rate or 5% loss in any individual step; time intervals are used when available. A load/strain shape check also requires loss per additional strain to exceed the advancing local baseline by more than six times, so accelerated progressive extension is not mistaken for sudden collapse; no additional strain supports collapse, while an unusable baseline requires review. Without a supported sudden event, use the measurement before first force below 10% of peak (ASTM E8/E8M-25 7.11.3.4). If neither crossing is recorded, continuous progressive unloading to at most 50% of peak with ongoing decline can support a review-flagged terminal estimate. Plateau/rising records and acquisition gaps do not get a final-reading fallback. Missing standards confirmation alone is recorded in Audit, not as a review failure. This hybrid and its safeguards are not a full standards-compliance claim. See guide for exact safeguards.',
        'instron_el_marker': 'Inspector-only: an Instron EL just outside CSV bounds within half the summary printed unit plus applicable CSV precision and floating-point allowance is displayed at the boundary. Tables and marker hover preserve the imported EL. This never changes analysis EL, reconstruction or specimen matching.',
        'failure_override': 'A specimen-wide manual EL override selects an original post-peak measurement row. It replaces the analysis endpoint in Calc EL, toughness, tensile curves and reconstruction, but not raw data, Instron results or pre-peak properties. Specimens identifies Manual versus Automatic selection; Audit retains the automatic result, selected row, reason and save time. Saved policy/history includes the source fingerprint and automatic snapshot. Stale overrides are not applied and require review.',
        'precision': 'Full numeric precision stored; two decimal places for ordinary results, five for elastic-fit R2, its warning threshold and fit/yield strain details. Counts and row identifiers remain integers.'}}
    from tensile_comparison import relative_comparison
    agreement = relative_comparison(frames)
    metadata['definitions']['relative_comparison'] = (
        'Calc-Instron agreement uses included specimens with finite paired measured values and nonzero Instron values. '
        'Mean and sample SD of 100*(Calc-Instron)/Instron are calculated per specimen, not from a ratio of group means. '
        'Positive means Calc is higher. Gauge-reconstructed values are not compared with uncorrected Instron. '
        'Failure EL is the selected measured CSV endpoint versus the Instron summary endpoint. '
        'SD Calc and SD Instron use exactly those same paired specimens in original units; '
        'SD change (%) is 100*(SD Calc/SD Instron-1), distinct from SD of paired percentage differences. '
        'SD change is unavailable for n<2 or zero Instron SD. Lower scatter is not proof of greater accuracy. '
        'Missing comparisons remain absent; n=1 has no SD. Checks are retained, not automatic exclusions. '
        'UTS is retained here for audit, but omitted from the comparison charts; discrepancies stay in specimen Checks.')
    sheets = {'Summary': summary, 'Specimens': specimens}
    if not agreement.empty:
        sheets['Calc-Instron agreement'] = agreement
    sheets.update({'Audit': checks, 'Export info': flatten_info(metadata)})
    write_workbook(path, sheets)


def export_curves(results, engine, path):
    sheets, seen, checks, names = {}, set(), [], {}
    def add(sheet, label, x, y, *, extra=None, info=None):
        x, y = np.asarray(x, float), np.asarray(y, float)
        signature = (sheet, label, x.tobytes(), y.tobytes(),
                     tuple((k, np.asarray(v, float).tobytes()) for k, v in (extra or {}).items()))
        if signature in seen:
            return
        seen.add(signature)
        key = (sheet, label)
        names[key] = names.get(key, 0) + 1
        if names[key] > 1:
            label += f' · variant {names[key]}'
        wh = sheet.startswith('WH')
        columns = sheets.setdefault(sheet, {})
        columns[label + (' · true plastic strain (%)' if wh else ' · strain (%)')] = pd.Series(x)
        y_label = (' · WH rate (MPa)' if wh else
                   ' · measured mean stress (MPa)' if sheet == 'Tensile pointwise' else ' · stress (MPa)')
        columns[label + y_label] = pd.Series(y)
        for suffix, values in (extra or {}).items():
            columns[label + ' · ' + suffix] = pd.Series(values)
        checks.append({'Sheet': sheet, 'Curve': label, 'Points': len(x),
                       'Start strain (%)': x[0] if len(x) else np.nan,
                       'End strain (%)': x[-1] if len(x) else np.nan, **(info or {})})

    for view_number, view in enumerate(results, 1):
        state, spec = view['state'], view['graph_spec']
        family = state['family']
        records, models = view.get('export_records', {}), view.get('export_models') or {}
        methods = {'landmark': ['landmark'], 'pointwise': ['pointwise'], 'comparison': ['pointwise', 'landmark']}.get(family, [])
        for group, model in models.items():
            basis = group_basis_label(state, spec, group)
            for method in methods:
                curve = model[method]
                if curve is None:
                    checks.append({'Sheet': 'Tensile ' + method, 'Curve': group, 'Status': model.get(method + '_diagnostic', 'Unavailable'), 'View': view_number})
                    continue
                x, y = curve['x'], curve['y']
                area = engine.compute_toughness_safe(x, y)
                specimen_areas = [engine.compute_toughness_safe(r['strain_pct'], r['stress_mpa']) for r in records[group]]
                mean = float(np.nanmean(specimen_areas))
                complete = method == 'landmark' or curve.get('complete', False)
                info = {'Group': group, 'Method': method, 'Basis': basis, 'View': view_number,
                        'Target gauge (mm)': group_policy(state, spec, group)['target_gauge_mm'] or np.nan,
                        'n': len(records[group]), 'Complete': complete,
                        'Curve toughness (MJ/m^3)': area, 'Mean specimen toughness (MJ/m^3)': mean,
                        'Difference (%)': 100 * (area - mean) / mean if complete and np.isfinite(mean) and mean else np.nan,
                        'Status': model.get(method + '_diagnostic', '')}
                extra = None
                if method == 'pointwise':
                    y = np.array(y, copy=True)
                    tail = np.array(y, copy=True)
                    split = len(curve['measured_x']) - 1
                    y[split + 1:] = np.nan
                    tail[:split] = np.nan
                    if not len(curve.get('predicted_x', [])):
                        tail[:] = np.nan
                    extra = {'projected stress (MPa)': tail}
                if method == 'landmark':
                    info['Shape n'] = curve.get('shape_n', curve.get('n'))
                    info['Property counts'] = json.dumps(curve.get('property_counts', {}), sort_keys=True)
                add('Tensile ' + method, group + ' · ' + basis, x, y, extra=extra, info=info)
                if method == 'landmark':
                    for point in curve.get('point_statistics', []):
                        entry = {'Sheet': 'Tensile landmark', 'Curve': group + ' · ' + basis,
                                 'Landmark': point['name'], 'Landmark strain (%)': point['x'],
                                 'Landmark stress (MPa)': point['y'], 'Strain SD (%)': point['x_sd'],
                                 'Stress SD (MPa)': point['y_sd'], 'Basis': basis,
                                 'Strain n': point.get('x_n', curve.get('n')),
                                 'Stress n': point.get('y_n', curve.get('n'))}
                        if entry not in checks:
                            checks.append(entry)
        if family == 'representative':
            for group, rows in records.items():
                basis = group_basis_label(state, spec, group)
                record = engine.choose_representative_curve(group, rows, spec.get('representative_overrides', {}).get(group))
                if record is not None:
                    add('Tensile representatives', group + ' · ' + record['sample'] + ' · ' + basis,
                        record['strain_pct'], record['stress_mpa'], info={'Basis': basis, 'Specimen ID': record['specimen_id'], 'View': view_number})
        if state['show_individuals'] and family in ('landmark', 'pointwise', 'comparison', 'representative'):
            for group, rows in records.items():
                basis = group_basis_label(state, spec, group)
                for record in rows:
                    add('Tensile individuals', group + ' · ' + record['sample'] + ' · ' + basis,
                        record['strain_pct'], record['stress_mpa'], info={'Basis': basis, 'Specimen ID': record['specimen_id'], 'View': view_number})
        for sheet, group, index, x, y in getattr(view['figure'], '_export_wh', []):
            label = group if index is None else group + ' · ' + records[group][index]['sample']
            info = {'Basis': 'Measured pre-peak response', 'View': view_number}
            if sheet == 'WH landmark':
                landmark = models.get(group, {}).get('landmark') or {}
                info.update({'Shape n': landmark.get('shape_n'),
                             'Property counts': json.dumps(landmark.get('property_counts', {}), sort_keys=True)})
            add(sheet, label, x, y, info=info)
    if not sheets:
        return  # scatter plots already have their data in the results workbook
    frames = {name: pd.DataFrame(columns) for name, columns in sheets.items()}
    frames['Curve checks'] = pd.DataFrame(checks)
    write_workbook(path, frames)
