"""Paired percentage differences; no research files or batch plotting."""
from contextlib import redirect_stdout
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
import ipywidgets as widgets

from selection_fixture import make_fixture
from tensile_comparison import relative_comparison, comparison_figure, comparison_properties, spread_table, ComparisonView
from tensile_workbench import TensileWorkbench


def frames():
    comparisons = pd.DataFrame([
        {'Included': included, 'Group': group, 'Specimen ID': ident, 'Property': prop,
         'Unit': unit, 'Calculated': calc, 'Instron': instron, 'Checks': check}
        for included, group, ident, prop, unit, calc, instron, check in [
            (True, 'A', 'a1', '0.2% yield strength', 'MPa', 110., 100., ''),
            (True, 'A', 'a2', '0.2% yield strength', 'MPa', 160., 200., 'Review identity'),
            (False, 'A', 'a3', '0.2% yield strength', 'MPa', 1000., 100., ''),
            (True, 'A', 'a4', '0.2% yield strength', 'MPa', 500., 0., ''),
            (True, 'A', 'a5', '0.2% yield strength', 'MPa', np.nan, 100., ''),
            (True, 'A', 'a6', '0.2% yield strength', 'MPa', 200., np.inf, ''),
            (True, 'B', 'b1', 'UTS', 'MPa', 700., 700., ''),
            (True, 'B', 'b1', 'Fitted elastic modulus', 'GPa', 100., np.nan, ''),
            (True, 'B', 'b1', 'Uniform elongation', '%', 10., 10., ''),
            (True, 'B', 'b2', 'Uniform elongation', '%', 11., 10., ''),
        ]])
    samples = pd.DataFrame([{'Included': True, 'Group': 'A', 'Specimen ID': 'a1',
        'CSV-derived fracture EL (%)': 5., 'Instron summary EL (%)': 4.,
        'Failure Elongation (%)': 70., 'Reconstructed EL (%)': 70., 'Checks': ''}])
    return {'instron_comparison': comparisons, 'tensile_samples': samples}


class RelativeComparisonTests(unittest.TestCase):
    def test_paired_percentage_mean_not_ratio_of_means_with_sample_sd(self):
        summary = relative_comparison(frames()).set_index(['Group', 'Property'])
        row = summary.loc[('A', '0.2% yield strength')]
        self.assertEqual(row['Paired n'], 2)
        self.assertAlmostEqual(row['Mean difference (%)'], -5.)
        self.assertAlmostEqual(row['SD difference (%)'], np.std([10., -20.], ddof=1))
        self.assertEqual(row['Mean Calc'], 135.)
        self.assertEqual(row['Mean Instron'], 150.)
        self.assertEqual(row['Pairs with checks'], 1)
        self.assertAlmostEqual(row['SD Calc'], np.std([110., 160.], ddof=1))
        self.assertAlmostEqual(row['SD Instron'], np.std([100., 200.], ddof=1))
        self.assertAlmostEqual(row['SD change (%)'], -50.)
        self.assertNotIn(('B', 'Fitted elastic modulus'), summary.index)
        self.assertTrue(np.isnan(summary.loc[('B', 'UTS'), 'SD difference (%)']))
        self.assertTrue(np.isnan(summary.loc[('B', 'UTS'), 'SD Calc']))
        self.assertEqual(summary.loc[('B', 'Uniform elongation'), 'SD Instron'], 0.)
        self.assertTrue(np.isnan(summary.loc[('B', 'Uniform elongation'), 'SD change (%)']))

    def test_measured_failure_elongation_only_and_no_data_is_empty(self):
        summary = relative_comparison(frames())
        el = summary[summary['Property'].eq('Failure elongation')].iloc[0]
        self.assertEqual(el['Mean difference (%)'], 25.)
        self.assertEqual(el['Mean Calc'], 5.)
        self.assertEqual(el['Paired n'], 1)
        self.assertTrue(np.isnan(el['SD change (%)']))
        self.assertTrue(relative_comparison({}).empty)

    def test_plot_traces_zero_reference_sd_and_pair_counts(self):
        summary = relative_comparison(frames())
        fig = comparison_figure(summary[summary.Property.eq('0.2% yield strength')])
        self.assertTrue(fig.layout.yaxis.zeroline)
        self.assertEqual(fig.layout.yaxis.rangemode, 'tozero')
        traces = {trace.meta['property']: trace for trace in fig.data}
        a = traces['0.2% yield strength']
        ys = list(a.x).index('A')
        self.assertEqual(a.y[ys], -5.)
        self.assertEqual(a.customdata[ys][0], 2)
        self.assertAlmostEqual(a.error_y.array[ys], np.std([10., -20.], ddof=1))
        self.assertEqual(fig.layout.height, 280)
        self.assertNotIn('UTS', comparison_properties(summary))
        with self.assertRaises(ValueError):
            comparison_figure(summary[summary.Property.eq('UTS')])
        for axis in fig.select_yaxes():
            self.assertIsNone(axis.matches)
            if axis.visible is not False:
                self.assertTrue(axis.zeroline)
                self.assertEqual(axis.rangemode, 'tozero')
        el = comparison_figure(summary[summary.Property.eq('Failure elongation')])
        self.assertIsNone(el.data[0].error_y.array[0])
        self.assertIsNone(el.layout.yaxis.matches)
        self.assertEqual(el.data[0].y[0], 25.)

    def test_spread_table_same_pairs_units_escaping_and_undefined_change(self):
        summary = relative_comparison(frames())
        ys = summary[summary.Property.eq('0.2% yield strength')].copy()
        ys['Group'] = '<script>not markup</script>'
        text = spread_table(ys)
        self.assertIn('SD (MPa)', text)
        self.assertIn('35.36', text)
        self.assertIn('70.71', text)
        self.assertIn('-50.00', text)
        self.assertIn('&lt;script&gt;', text)
        self.assertNotIn('<script>', text)
        single = spread_table(summary[summary.Property.eq('Failure elongation')])
        self.assertEqual(single.count('—'), 3)
        zero_sd = spread_table(summary[summary.Property.eq('Uniform elongation')])
        self.assertIn('0.00', zero_sd)
        self.assertEqual(zero_sd.count('—'), 1)

    def test_resize_preserves_independent_panels_and_values(self):
        view = ComparisonView(widgets)
        self.addCleanup(view.clear)
        view.set_frames(frames(), ['A', 'B'])
        view.ui.selected_index = 0
        self.assertEqual(view._columns, 3)
        self.assertNotIn('UTS', view.charts)
        before = [(prop, tuple(chart.data[0].y)) for prop, chart in view.charts.items()]
        for card in view.grid.children:
            self.assertIn('Spread · SD', card.children[1].value)
        view._resize(600)
        self.assertEqual(view._columns, 1)
        self.assertTrue(all(chart.layout.width == 598 for chart in view.charts.values()))
        view._resize(1000)
        self.assertEqual(view._columns, 2)
        view._resize(1200)
        self.assertEqual(view._columns, 3)
        self.assertEqual(view.grid.layout.grid_template_columns, 'repeat(3, minmax(0, 1fr))')
        after = [(prop, tuple(chart.data[0].y)) for prop, chart in view.charts.items()]
        self.assertEqual(before, after)
        self.assertTrue(all(chart.layout.yaxis.matches is None for chart in view.charts.values()))

    def test_chart_is_lazy_reuses_tables_and_clears_when_data_removed(self):
        view = ComparisonView(widgets)
        self.addCleanup(view.clear)
        with patch('tensile_comparison.comparison_figure', wraps=comparison_figure) as build:
            view.set_frames(frames(), ['A'])
            self.assertFalse(view.charts)
            build.assert_not_called()
            view.ui.selected_index = 0
            self.assertTrue(view.charts)
            self.assertEqual(build.call_count, 2)
            self.assertEqual(set(view.summary['Group']), {'A'})
            view.set_frames(frames(), ['A'])
            self.assertEqual(build.call_count, 2)
            view.ui.selected_index = None
            view.ui.selected_index = 0
            self.assertEqual(build.call_count, 2)
            view.set_frames({}, [])
            self.assertFalse(view.charts)
            self.assertIn('No included specimens', view.box.children[0].value)

    def test_uts_only_has_no_chart_but_mismatch_still_flagged(self):
        from tensile_instron import _verify_reference
        from tensile_tables import specimen_check_failures
        reference = _verify_reference({'values': {'uts': 700.}, 'resolution': {'uts': .01}},
            {'_acquisition': {'stress': [0., 701., 300.], 'check_resolution': {'uts': .01}}}, True)
        self.assertIn('UTS: CSV maximum 701.00 MPa vs Instron 700.00 MPa', '\n'.join(reference['check_failures']))
        self.assertIn('UTS:', specimen_check_failures({}, reference))
        view = ComparisonView(widgets)
        self.addCleanup(view.clear)
        data = frames()['instron_comparison']
        view.set_frames({'instron_comparison': data[data.Property.eq('UTS')]}, ['B'])
        view.ui.selected_index = 0
        self.assertFalse(view.charts)
        self.assertIn('UTS discrepancies', view.box.children[0].value)

    def test_integration_below_summary_and_export_agrees(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory(prefix='agreement-export-synthetic-') as temp:
            root = Path(temp)
            make_fixture(root)
            with redirect_stdout(io.StringIO()):
                app = TensileWorkbench(root)
            raw_match = app.session.instron.match
            def with_values(group, record):
                reference = raw_match(group, record)
                p = app.session.engine.specimen_properties(record.get('_measured_record', record))
                return {**reference, 'values': {**reference['values'], 'ys': p['Yield (MPa)'] / 1.1,
                                                'el': p['Failure elongation (%)'] / 1.05}}
            with patch.object(app.session.instron, 'match', side_effect=with_values):
                app.session._property_frames.clear()
                app._refresh_properties()
                self.assertIs(app.tables.tabs.children[0].children[1], app.tables.comparison.ui)
                expected = app.tables.comparison.summary
                self.assertFalse(expected.empty)
                self.assertFalse(app.tables.comparison.charts)
                folder = app.session.export_properties(app.state(), app._current()['definition'])
            workbook = load_workbook(next(folder.glob('*_results.xlsx')), data_only=True)
            self.assertIn('Calc-Instron agreement', workbook.sheetnames)
            values = list(workbook['Calc-Instron agreement'].values)
            result = pd.DataFrame(values[1:], columns=values[0])
            np.testing.assert_allclose(result['Mean difference (%)'], expected['Mean difference (%)'])
            self.assertEqual(list(result['Paired n']), list(expected['Paired n']))
            for column in ['SD Calc', 'SD Instron', 'SD change (%)']:
                np.testing.assert_allclose(pd.to_numeric(result[column]), expected[column], equal_nan=True)
            workbook.close()
            app.tables.inspector.clear()
            app.tables.comparison.clear()


if __name__ == '__main__':
    unittest.main()
