"""Permanent variables, reload checks and split specimen charts; synthetic data only."""
from contextlib import redirect_stdout
from copy import deepcopy
import io
import json
from pathlib import Path
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
import test_group_properties as fixture
from tensile_locations import (LOCATION, SPECIMEN_COUNT, LEGACY_NUMBER,
                               ensure_required_variables, reconcile_specimen_counts)
from tensile_property_variants import (VARIANTS_KEY, SPECIMEN_VARIANTS_KEY,
                                       SPECIMEN_FAMILY, split_property_charts)
from tensile_moves import migrate_settings, forget_group, references
from tensile_variables import validate_variables, variable_usage
from tensile_workbench import TensileWorkbench
from tensile_plotly import figure_to_plotly
from variable_fixture import column, cell, set_cell


class RequiredLocationTests(unittest.TestCase):
    setUp = fixture.GroupPropertiesTests.setUp

    def test_count_startup_reload_and_no_lowering_reserved_total(self):
        app = self.app
        self.assertEqual(app.session.imported_counts, {'Alloy A': 3, 'Alloy Z': 3, 'Heat treated': 3})
        view = app.group_variables
        self.assertEqual(view.table.columns[0]['id'], SPECIMEN_COUNT)
        self.assertEqual(cell(view, SPECIMEN_COUNT, 'Alloy A'), '3')
        self.assertEqual(view.table.defaults[SPECIMEN_COUNT]['Alloy A'], '3')
        # Too-small saves are raised, while a larger planned total is retained.
        set_cell(view, SPECIMEN_COUNT, 'Alloy A', '1')
        view.save_all.click()
        self.assertEqual(cell(view, SPECIMEN_COUNT, 'Alloy A'), '3')
        set_cell(view, SPECIMEN_COUNT, 'Alloy A', '10')
        view.save_all.click()
        self.assertEqual(cell(view, SPECIMEN_COUNT, 'Alloy A'), '10')
        self.assertEqual(view.table.defaults[SPECIMEN_COUNT]['Alloy A'], '3')
        project = deepcopy(app.store.data)
        project['specimen_exclusions'] = {'Alloy A/coupon_1.csv': 'excluded still counts'}
        project['group_variables'][SPECIMEN_COUNT]['values']['Alloy Z'] = 0
        app.store.save(project)
        with redirect_stdout(io.StringIO()), patch.object(TensileWorkbench, '_refresh_properties'):
            reopened = TensileWorkbench(self.root)
        totals = reopened.store.data['group_variables'][SPECIMEN_COUNT]['values']
        self.assertEqual(totals['Alloy A'], 10)
        self.assertEqual(totals['Alloy Z'], 3)
        reopened.group_variables.all_groups.value = True
        self.assertEqual(reopened.group_variables.table.defaults[SPECIMEN_COUNT]['Alloy A'], '3')
        self.assertEqual(cell(reopened.group_variables, SPECIMEN_COUNT, 'Alloy A'), '10')
        new = self.root / 'data' / 'Alloy Z' / 'coupon_4.csv'
        new.write_bytes((new.parent / 'coupon_1.csv').read_bytes())
        with (patch.object(reopened.session, '_load', side_effect=AssertionError('No curve analysis to count')),
              patch.object(reopened, '_refresh_properties')):
            reopened._reload_data(plots=False)
        self.assertNotIn('failed', reopened.status.value.lower())
        self.assertEqual(reopened.store.data['group_variables'][SPECIMEN_COUNT]['values']['Alloy Z'], 4)
        self.assertEqual(reopened.store.data['group_variables'][SPECIMEN_COUNT]['values']['Alloy A'], 10)
        self.assertEqual(reopened.group_variables.table.defaults[SPECIMEN_COUNT]['Alloy Z'], '4')
        self.assertEqual(cell(reopened.group_variables, SPECIMEN_COUNT, 'Alloy Z'), '4')
        new.unlink()
        with patch.object(reopened, '_refresh_properties'):
            reopened._reload_data(plots=False)
        self.assertEqual(reopened.store.data['group_variables'][SPECIMEN_COUNT]['values']['Alloy Z'], 4)
        self.assertEqual(reopened.group_variables.table.defaults[SPECIMEN_COUNT]['Alloy Z'], '3')
        self.assertEqual(cell(reopened.group_variables, SPECIMEN_COUNT, 'Alloy Z'), '4')

    def test_required_variables_are_integer_locked_and_not_deletable(self):
        app = self.app
        app.custom_variables_panel.selected_index = 0
        app.custom_variable_tabs.selected_index = 1
        for scope, view, ident in (('group', app.group_variables, SPECIMEN_COUNT),
                                    ('specimen', app.specimen_variables, LOCATION)):
            self.assertEqual(view.table.columns[0]['id'], ident)
            self.assertIn(ident, view.table.required)
            for removed in ('view', 'single_ui', 'choice', 'fields', 'save', 'remove'):
                self.assertFalse(hasattr(view, removed), removed)
            self.assertNotEqual(view.table_ui.layout.display, 'none')
            view._request_delete(ident)
            self.assertIsNone(view._pending_delete)
            before = deepcopy(view.table.columns)
            view.table._message(None, dict(type='delete', id=ident), [])
            self.assertEqual(view.table.columns, before)
            view.table._message(None, dict(type='paste', row=0, column=0, text='Rename\nmm'), [])
            self.assertIn('fixed', view.table.notice)
            self.assertEqual(view.table.columns, before)
            view.table._message(None, dict(type='paste', row=2, column=0, text='1.5'), [])
            self.assertIn('whole numbers', view.table.notice)
            self.assertEqual(view.table.columns, before)
            variables = deepcopy(app.store.data[scope + '_variables'])
            variables.pop(ident)
            with self.assertRaisesRegex(ValueError, 'cannot be deleted'):
                app._save_variables(variables, scope)
            variable = deepcopy(app.store.data[scope + '_variables'][ident])
            variable['values'][next(iter(view.table.groups))] = 1.5
            with self.assertRaisesRegex(ValueError, 'whole numbers'):
                validate_variables({ident: variable}, scope=scope)

    def test_default_location_not_saved_as_override_and_blank_restores_it(self):
        for record, number in zip(self.session._load('Alloy A'), (1, 4, 7)):
            record['specimen_label'] = f'Wall_{number:02}'
        app = self.app
        app.custom_variables_panel.selected_index = 0
        app.custom_variable_tabs.selected_index = 1
        view = app.specimen_variables
        self.assertEqual([cell(view, LOCATION, g) for g in view.table.groups], ['1', '4', '7'])
        view.save_all.click()
        self.assertEqual(app.store.data['specimen_variables'][LOCATION]['values'], {})
        set_cell(view, LOCATION, 'Alloy A/coupon_2.csv', '9')
        view.save_all.click()
        self.assertEqual(app.store.data['specimen_variables'][LOCATION]['values'], {'Alloy A/coupon_2.csv': 9})
        view.table._message(None, dict(type='paste', row=3, column=0, text=''), [])
        view.save_all.click()
        self.assertEqual(app.store.data['specimen_variables'][LOCATION]['values'], {})
        self.assertEqual(view.table.defaults[LOCATION]['Alloy A/coupon_2.csv'], '4')
        self.assertEqual(cell(view, LOCATION, 'Alloy A/coupon_2.csv'), '4')

    def test_count_context_is_not_editable_or_pasted_and_unknown_count_stays_unknown(self):
        view = self.app.group_variables
        view.table._message(None, dict(type='paste', row=2, column=0, text='10'), [])
        view.save_all.click()
        self.assertEqual(view.table.defaults[SPECIMEN_COUNT]['Alloy A'], '3')
        self.assertEqual(column(view, SPECIMEN_COUNT)['values']['Alloy A'], '10')
        self.assertEqual(self.app.store.data['group_variables'][SPECIMEN_COUNT]['values']['Alloy A'], 10)
        before = deepcopy(view.table.columns)
        view.table._message(None, dict(type='paste', row=2, column=0, text='3 tested of 10'), [])
        self.assertIn('Nothing pasted', view.table.notice)
        self.assertEqual(view.table.columns, before)
        # A saved, unavailable group is not assumed to have zero tested specimens.
        self.app._save_group_variable(SPECIMEN_COUNT, dict(name='Specimen count', unit='', values={'Missing group': 12}))
        view.refresh(['Alloy A'], force=True)
        view.all_groups.value = True
        self.assertIn('Missing group', view.table.groups)
        self.assertNotIn('Missing group', view.table.defaults[SPECIMEN_COUNT])
        self.assertEqual(cell(view, SPECIMEN_COUNT, 'Missing group'), '12')

    def test_migrate_trial_location_and_variants_without_duplicate_axes(self):
        project = deepcopy(self.app.store.data)
        project['specimen_variables'] = {'old-location': dict(name='Location', unit='', values={'Alloy A/coupon_1.csv': 2})}
        settings = project['graphs'][0]['settings']
        settings.pop(SPECIMEN_VARIANTS_KEY, None)
        settings['families'] = ['properties_by_group']
        settings[VARIANTS_KEY] = [
            dict(id='group', name='Gas series', settings={'properties_by_group_level': 'group'}),
            dict(id='individual', name='Positions', settings={'properties_by_group_level': 'specimen',
                 'properties_by_group_x_variable': 'old-location', 'properties_by_group_planned_positions': '1:10'})]
        migrated = split_property_charts(ensure_required_variables(project))
        self.assertEqual(set(migrated['specimen_variables']), {LOCATION})
        self.assertEqual(migrated['specimen_variables'][LOCATION]['values']['Alloy A/coupon_1.csv'], 2)
        settings = migrated['graphs'][0]['settings']
        self.assertEqual([v['id'] for v in settings[VARIANTS_KEY]], ['group'])
        self.assertEqual([v['id'] for v in settings[SPECIMEN_VARIANTS_KEY]], ['individual'])
        self.assertEqual(settings['families'], ['properties_by_group', SPECIMEN_FAMILY])
        self.assertEqual(settings[SPECIMEN_VARIANTS_KEY][0]['settings']['properties_by_group_x_variable'], LOCATION)
        self.assertEqual(migrated['group_variables'][SPECIMEN_COUNT]['values']['Alloy A'], 10)
        self.assertEqual(split_property_charts(ensure_required_variables(migrated)), migrated)
        self.assertTrue(variable_usage(migrated, LOCATION, scope='specimen'))
        # Fractional custom locations keep their data and references, not rounded away.
        project['specimen_variables']['old-location']['values']['Alloy A/coupon_1.csv'] = 2.5
        legacy = ensure_required_variables(project)
        self.assertEqual(legacy['specimen_variables']['old-location']['values']['Alloy A/coupon_1.csv'], 2.5)
        self.assertEqual(legacy['specimen_variables']['old-location']['name'], 'Location (legacy)')
        self.assertEqual(legacy['specimen_variables'][LOCATION]['values'], {})

    def test_counts_move_without_conflict_with_discovered_minimum_and_forget(self):
        project = reconcile_specimen_counts(self.app.store.data, {'Owner/Alloy A': 3})
        project['group_variables'][SPECIMEN_COUNT]['values']['Alloy A'] = 10
        project['specimen_variables'][LOCATION]['values']['Alloy A/coupon_1.csv'] = 5
        moved = migrate_settings(project, dict(old='Alloy A', new='Owner/Alloy A',
            mapping={'Alloy A/coupon_1.csv': 'Owner/Alloy A/coupon_1.csv'}), audit=False)
        self.assertEqual(moved['group_variables'][SPECIMEN_COUNT]['values']['Owner/Alloy A'], 10)
        self.assertEqual(moved['specimen_variables'][LOCATION]['values']['Owner/Alloy A/coupon_1.csv'], 5)
        forgotten = forget_group(moved, 'Owner/Alloy A')
        for key in ('values', 'minimums'):
            self.assertNotIn('Owner/Alloy A', forgotten['group_variables'][SPECIMEN_COUNT][key])
        # Automatic inventory counts alone must not trigger hashing every group for moves.
        groups, _ = references(self.app.store.data)
        self.assertNotIn('Heat treated', groups)
        self.assertNotIn('Alloy Z', groups)

    def test_split_exports_include_all_groups_and_count_invalidates_preview(self):
        for group in self.session.files:
            for record, number in zip(self.session._load(group), (1, 4, 7)):
                record['specimen_label'] = f'Wall_{number:02}'
        app = self.app
        app.controls['groups'].value = ('Alloy A', 'Alloy Z', 'Heat treated')
        app._display_order = ['Heat treated', 'Alloy A', 'Alloy Z']
        app.family_boxes[SPECIMEN_FAMILY].value = True
        with patch.object(app.viewer, 'set_items') as shown:
            app.update()
        self.assertEqual(len(app._results), 3, app.details.value)
        self.assertEqual([r['state']['groups'] for r in app._results], [['Heat treated'], ['Alloy A'], ['Alloy Z']])
        self.assertEqual(len({i['key'] for i in shown.call_args.args[0]}), 3)
        self.session.engine.PLOT_DPI = 30
        destination = self.session.export(app._results, include_tables=True)
        self.assertEqual(len(list(destination.glob('*.png'))), 3)
        workbook = load_workbook(next(destination.glob('*_results.xlsx')), read_only=True, data_only=True)
        rows = list(workbook['Group variables'].values)
        self.assertEqual({row[0] for row in rows[1:]}, {'Alloy A', 'Alloy Z', 'Heat treated'})
        for sheet_name, ident in (('Group variables', SPECIMEN_COUNT), ('Specimen variables', LOCATION)):
            sheet = workbook[sheet_name]
            header = list(next(sheet.values))
            variable_col, value_col = header.index('Variable ID'), header.index('Value')
            cells = [row[value_col] for row in sheet.iter_rows(min_row=2) if row[variable_col].value == ident]
            self.assertTrue(cells)
            self.assertTrue(all(cell.number_format == '0' for cell in cells))
        workbook.close()
        metadata = json.loads((destination / 'settings.json').read_text())
        self.assertEqual(len(metadata['sources']), 9)
        old = app._results[1]
        app._save_group_variable(SPECIMEN_COUNT, dict(name='Specimen count', unit='', values={'Alloy A': 10}))
        with self.assertRaisesRegex(ValueError, 'variables changed'):
            self.session.export([old])
        updated = self.session.render(old['state'])
        self.assertIsNot(old['figure'], updated['figure'])
        chart = figure_to_plotly(updated['figure'])
        self.assertEqual(list(chart.layout.xaxis.tickvals), list(range(1, 11)))
        self.assertEqual(chart.layout.xaxis.tickformat, '.0f')
        self.assertEqual(chart.layout.xaxis.hoverformat, '.0f')


if __name__ == '__main__':
    unittest.main()
