"""Range-sensitivity warnings must be visible without changing fitted results."""
from copy import deepcopy
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

import ipywidgets as w
import numpy as np
from openpyxl import load_workbook

import tensile_core as core
from tensile_exports import export_results
from tensile_inspector import inspection_checks, inspection_summary
from tensile_properties import ELASTIC_SENSITIVITY_DEFAULTS, specimen_calculation
from tensile_tables import property_tables, specimen_review_mask, specimen_check_failures
from tensile_yield_review import GroupYieldReview


def sensitive_record():
    # Two individually plausible slopes before the main yield transition.
    # Their disagreement warrants review, not a claim that either is correct.
    stress = np.linspace(0., 1000., 8000)
    strain = 100 * (np.minimum(stress, 300.) / 120000.
                    + np.maximum(stress - 300., 0.) / 100000.
                    + np.maximum(stress - 900., 0.) / 2000.)
    return {'sample': 'coupon_1', 'source_file': 'synthetic/coupon_1.csv',
            'specimen_id': 'Synthetic/coupon_1.csv', 'source_sha256': 'a' * 64,
            'strain_pct': np.r_[strain, strain[-1] + [.1, .11]],
            'stress_mpa': np.r_[stress, 500., 0.]}


def no_reference():
    return {'values': {}, 'status': 'Unavailable', 'source': '', 'notes': '',
            'label': '', 'index': None, 'excluded': False, 'sha256': ''}


class FitSensitivityTests(unittest.TestCase):
    def test_warning_retains_exact_fit_mask_yield_and_other_properties(self):
        record = sensitive_record()
        before = deepcopy(record)
        with patch('tensile_properties._elastic_fit_sensitivity', return_value=dict(ELASTIC_SENSITIVITY_DEFAULTS)):
            baseline = specimen_calculation(record)
        actual = specimen_calculation(record)
        p = actual['properties']
        self.assertTrue(p['Fit review required'])
        self.assertIn('sensitive to selected range', p['Automatic fit sensitivity warning'])
        self.assertGreater(p['Automatic fit sensitivity (%)'], 10.)
        self.assertGreater(p['Elastic fit R2'], .999)
        self.assertEqual(p['Yield status'], 'resolved')
        self.assertEqual(p['Automatic fit warning'], '')  # fatal-fit failures are separate
        for field in ('Fitted E (GPa)', 'Yield (MPa)', 'Yield strain (%)', 'Elastic fit R2',
                      'UTS (MPa)', 'Failure elongation (%)', 'Toughness (MJ/m^3)'):
            self.assertEqual(p[field], baseline['properties'][field])
        np.testing.assert_array_equal(actual['elastic_mask'], baseline['elastic_mask'])
        for key in ('strain_pct', 'stress_mpa'):
            np.testing.assert_array_equal(record[key], before[key])

    def test_manual_review_suppresses_active_warning_but_retains_automatic_audit(self):
        record = sensitive_record()
        saved = {'mode': 'range', 'strain_bounds': [.1, .2], 'source_sha256': 'a' * 64}
        automatic = specimen_calculation(record)['properties']
        p = specimen_calculation(record, override=saved)['properties']
        self.assertEqual(p['Fit method'], 'Manual range')
        self.assertFalse(p['Fit review required'])
        self.assertEqual(p['Automatic fit sensitivity warning'], automatic['Automatic fit sensitivity warning'])
        self.assertNotIn('sensitive to selected range', specimen_check_failures(p, no_reference()))
        self.assertAlmostEqual(p['Fitted E (GPa)'], 120.)
        stale = specimen_calculation(record, override={**saved, 'source_sha256': 'b' * 64})['properties']
        self.assertTrue(stale['Fit review required'])
        self.assertIn('sensitive to selected range', specimen_check_failures(stale, no_reference()))

    def test_warnings_follow_property_eligibility_and_inspector_uses_same_message(self):
        record = sensitive_record()
        calculation = specimen_calculation(record)
        p, reference = calculation['properties'], no_reference()
        for mode in ('all', 'exclude_shape', 'after_uts', 'after_yield'):
            self.assertIn(p['Automatic fit sensitivity warning'], specimen_check_failures(p, reference, mode=mode))
        for mode in ('no_strain', 'exclude'):
            self.assertNotIn('sensitive to selected range', specimen_check_failures(p, reference, mode=mode))
        payload = dict(record=record, calculation=calculation, reference=reference,
                       sample=record['sample'], group='Synthetic', included=True,
                       source_file=record['source_file'], exclusion_reason='')
        self.assertIn(p['Automatic fit sensitivity warning'], inspection_checks(payload))
        self.assertIn('another supported interval', inspection_summary(payload))
        self.assertIn('not a confidence interval', inspection_summary(payload))

    def test_group_overview_identifies_sensitive_fit_and_counts_it_for_review(self):
        record = sensitive_record()
        review = GroupYieldReview(w, Mock(return_value=[record]))
        self.addCleanup(review.clear)
        review.set_rows([{'id': record['specimen_id'], 'group': 'Synthetic', 'sample': record['sample'],
                          'included': True, 'mode': 'all'}])
        review.set_active(True)
        self.assertIn('Fit sensitive to selected range', review.grid.children[0].children[0].value)
        self.assertIn('1 fit needs review', review.status.value)
        self.assertIn('Review fit', review.entries[0]['figure'].layout.annotations[-1].text)

    def test_warning_and_numeric_diagnostics_survive_tables_and_excel_export(self):
        record = sensitive_record()
        frames = property_tables({'Synthetic': [record]}, {}, core,
                                 SimpleNamespace(match=lambda group, row: no_reference()))
        samples = frames['tensile_samples']
        self.assertTrue(specimen_review_mask(samples).iloc[0])
        self.assertTrue(samples.iloc[0]['Included'])
        self.assertIn('sensitive to selected range', samples.iloc[0]['Checks'])
        fields = ('Automatic fit sensitivity (%)', 'Automatic fit sensitivity alternative E (GPa)',
                  'Automatic fit sensitivity alternative lower (% UTS)',
                  'Automatic fit sensitivity alternative upper (% UTS)',
                  'Automatic fit sensitivity alternative R2')
        audit = frames['specimen_diagnostics'].iloc[0]
        self.assertTrue(all(np.isfinite(audit[field]) for field in fields))
        with tempfile.TemporaryDirectory(prefix='fit-sensitivity-synthetic-') as directory:
            path = Path(directory) / 'synthetic.xlsx'
            export_results(frames, path, {})
            book = load_workbook(path)
            try:
                sheet = book['Audit']
                headers = [cell.value for cell in sheet[1]]
                for field in fields:
                    self.assertAlmostEqual(sheet.cell(2, headers.index(field) + 1).value, audit[field])
                specimen_headers = [cell.value for cell in book['Specimens'][1]]
                checks = book['Specimens'].cell(2, specimen_headers.index('Checks') + 1).value
                self.assertIn('sensitive to selected range', checks)
            finally:
                book.close()


if __name__ == '__main__':
    unittest.main()
