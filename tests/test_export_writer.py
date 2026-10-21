"""Excel writer regression checks; only small temporary synthetic workbooks."""
from datetime import datetime, timedelta
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from xml.etree import ElementTree
from zipfile import ZipFile

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import IllegalCharacterError
from openpyxl.worksheet._writer import ALL_TEMP_FILES

from tensile_exports import write_workbook
from tensile_fit import fit_display_places


class ExportWriterTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory(prefix='excel-writer-synthetic-')
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.path = self.root / 'synthetic.xlsx'

    def open_book(self):
        book = load_workbook(self.path)
        self.addCleanup(book.close)
        return book

    def test_values_precision_literal_text_and_layout(self):
        frame = pd.DataFrame({
            'Specimen': ['=1+1', '#N/A'], 'Checks': ['First\nSecond', ''],
            'UTS (MPa)': [800.123456789, 700.987654321],
            'Elastic fit R²': [.987654321, .999876543],
            'Fit lower strain (%)': [.0123456789, .02456789],
            'Included n': [np.int64(3), np.int64(2)],
            'Failure row (1-based)': [100, 200], 'Included': [True, False],
            'Metadata': [{'range': [1, 2]}, (3, 4)],
            '=Literal header': ['ordinary', 'text'],
        })
        write_workbook(self.path, {'Summary': frame, 'Curve checks': frame.iloc[:1]})
        book = self.open_book()
        self.assertEqual(book.sheetnames, ['Summary', 'Curve checks'])
        sheet = book['Summary']
        self.assertEqual([cell.value for cell in sheet[1]], list(frame.columns))
        self.assertEqual(sheet.freeze_panes, 'B2')
        self.assertEqual(book['Curve checks'].freeze_panes, 'A2')
        self.assertFalse(sheet.sheet_view.showGridLines)
        self.assertEqual(sheet.auto_filter.ref, 'A1:J3')
        self.assertEqual(sheet.row_dimensions[1].height, 45)
        self.assertEqual(sheet.column_dimensions['B'].width, 65)
        self.assertTrue(sheet['A1'].font.bold)
        self.assertEqual(sheet['A1'].fill.fgColor.rgb, '0024445C')
        self.assertTrue(sheet['A1'].alignment.wrap_text)
        self.assertEqual(sheet['A2'].font.name, 'Arial')
        self.assertTrue(sheet['B2'].alignment.wrap_text)
        for address in ('A2', 'A3', 'J1'):
            self.assertEqual(sheet[address].data_type, 's')
        self.assertEqual(sheet['A2'].value, '=1+1')
        self.assertEqual(sheet['A3'].value, '#N/A')
        self.assertEqual(sheet['C2'].value, frame.iloc[0, 2])
        self.assertEqual(sheet['C2'].number_format, '0.00')
        self.assertEqual(sheet['D2'].value, frame.iloc[0, 3])
        self.assertEqual(sheet['D2'].number_format, '0.00000')
        self.assertEqual(sheet['E2'].number_format, '0.00000')
        self.assertEqual(sheet['F2'].number_format, '0')
        self.assertEqual(sheet['G2'].number_format, '0')
        self.assertEqual(sheet['H2'].data_type, 'b')
        self.assertEqual(json.loads(sheet['I2'].value), {'range': [1, 2]})
        self.assertEqual(json.loads(sheet['I3'].value), [3, 4])

    def test_ragged_curves_keep_every_point_without_styled_blank_padding(self):
        strain = np.linspace(.123456789, 9.987654321, 127)
        stress = strain * 71.23456789
        frame = pd.DataFrame({'A strain': pd.Series(strain), 'A stress': pd.Series(stress),
                              'B strain': pd.Series(strain[:9]), 'B stress': pd.Series(stress[:9])})
        write_workbook(self.path, {'Tensile individuals': frame})
        sheet = self.open_book()['Tensile individuals']
        np.testing.assert_allclose([sheet.cell(i + 2, 1).value for i in range(127)], strain, rtol=1e-14)
        np.testing.assert_allclose([sheet.cell(i + 2, 2).value for i in range(127)], stress, rtol=1e-14)
        self.assertEqual(sheet.auto_filter.ref, 'A1:D128')
        self.assertIsNone(sheet['C11'].value)
        with ZipFile(self.path) as archive:
            xml = ElementTree.fromstring(archive.read('xl/worksheets/sheet1.xml'))
            cells = xml.findall('.//{*}c')
            self.assertEqual(len(cells), 4 + 2 * 127 + 2 * 9)
            self.assertNotIn('C11', {cell.attrib['r'] for cell in cells})

    def test_missing_values_and_date_formats(self):
        frame = pd.DataFrame({
            'Missing': [None, np.nan, np.inf, -np.inf, pd.NA, pd.NaT],
            'Date': [datetime(2026, 10, 21)] * 6,
            'Duration': [timedelta(seconds=30)] * 6,
        })
        write_workbook(self.path, {'Specimens': frame})
        sheet = self.open_book()['Specimens']
        self.assertTrue(all(sheet.cell(row, 1).value is None for row in range(2, 8)))
        self.assertEqual(sheet['B2'].value, datetime(2026, 10, 21))
        self.assertEqual(sheet['B2'].number_format, 'yyyy-mm-dd h:mm:ss')
        self.assertEqual(sheet['C2'].value, timedelta(seconds=30))
        self.assertEqual(sheet['C2'].number_format, '[hh]:mm:ss')

    def test_empty_sheets_and_headers_only(self):
        write_workbook(self.path, {'Empty': pd.DataFrame(),
                                  'Audit': pd.DataFrame(columns=['Specimen', 'Elastic fit R2'])})
        book = self.open_book()
        self.assertEqual(book['Empty']['A1'].value, 'No data available')
        self.assertEqual(book['Audit'].auto_filter.ref, 'A1:B1')
        self.assertEqual(book['Audit'].freeze_panes, 'B2')

    def test_formats_are_computed_per_column_not_per_datapoint(self):
        frame = pd.DataFrame({'Stress': np.arange(1000.), 'R2': np.ones(1000), 'Included n': [2] * 1000})
        with patch('tensile_exports.fit_display_places', wraps=fit_display_places) as formats:
            write_workbook(self.path, {'Synthetic': frame})
        self.assertEqual(formats.call_count, 2)
        with ZipFile(self.path) as archive:
            styles = ElementTree.fromstring(archive.read('xl/styles.xml'))
            self.assertLessEqual(len(styles.find('{*}cellXfs')), 6)
            self.assertLessEqual(len(styles.find('{*}fonts')), 3)

    def test_failed_save_preserves_destination_and_cleans_temporary_files(self):
        frame = pd.DataFrame({'Value': [1.23456789, 2.3456789]})
        write_workbook(self.path, {'Original': frame})
        original = self.path.read_bytes()
        pending = set(ALL_TEMP_FILES)
        with patch('openpyxl.workbook.workbook.Workbook.save', side_effect=OSError('Synthetic save failure')):
            with self.assertRaisesRegex(OSError, 'Synthetic save failure'):
                write_workbook(self.path, {'Replacement': frame})
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(set(ALL_TEMP_FILES), pending)
        self.assertEqual(list(self.root.iterdir()), [self.path])

    def test_invalid_cell_also_cleans_streams(self):
        pending = set(ALL_TEMP_FILES)
        with self.assertRaises(IllegalCharacterError):
            write_workbook(self.path, {'Invalid': pd.DataFrame({'Text': ['Valid', '\x00']})})
        self.assertEqual(set(ALL_TEMP_FILES), pending)
        self.assertFalse(self.path.exists())


if __name__ == '__main__':
    unittest.main()
